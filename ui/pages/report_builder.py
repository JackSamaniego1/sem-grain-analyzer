"""
Reports page logic — no widgets (REP-05/06/07).

* :func:`collect_inputs` snapshots the open session's FILTERED results on the
  GUI thread; :func:`build_model` turns them into a ``reports.ReportModel``
  off the GUI thread (overlays are written to ``<session>/report_assets``).
* :func:`results_fingerprint` identifies the numbers a report was built from,
  so the page can tell when re-analysis / filter edits made it stale.
* :func:`merge_refresh` swaps fresh numbers into an edited report while
  keeping every user edit (title, captions, notes, order, sections, units,
  bins, grain annotations, export history).
* :func:`render_outputs` renders Excel and/or PowerPoint from a plain dict
  copy of the model (safe to run on a pool thread).

Grain exclusion is NOT stored in the report: it goes through the app's
manual-exclusion mechanism (``AppState.delete_grains``) so Review, report and
exports always agree.  Per-grain annotations live in the report's grain rows
(``grain["note"]``) and are appended to the image notes on export.
"""
from __future__ import annotations

import copy
import hashlib
import json
import os
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Sequence, Tuple

from reports.model import ImageSummary, ReportImageInput, ReportModel, Section

REPORT_ASSETS = "report_assets"
FIXED_ORDER = ("cover", "overview_table", "combined_distribution", "images", "parameters",
               "raw_data")
MAX_EXPORT_HISTORY = 25

# Colour of each Excel tab / outline chip.  Mirrors reports.charts.TAB_COLORS
# (the renderer's palette) so the outline reads like the workbook's tabs.
from reports.charts import TAB_COLORS  # noqa: E402

SECTION_COLORS = {
    "cover": TAB_COLORS["overview"], "overview_table": TAB_COLORS["overview"],
    "combined_distribution": TAB_COLORS["charts"], "image": TAB_COLORS["image"],
    "images": TAB_COLORS["image"], "parameters": TAB_COLORS["methods"],
    "raw_data": TAB_COLORS["raw"], "custom_text": "#6A4C93",
}

SECTION_LABELS = {
    "cover": "Cover", "overview_table": "Overview table",
    "combined_distribution": "Summary charts", "images": "Images",
    "parameters": "Methods", "raw_data": "Raw data", "custom_text": "Text",
}

# Where each section ends up in the two outputs (shown in the preview header).
SECTION_TARGETS = {
    "cover": ("Overview sheet header", "Title slide"),
    "overview_table": ("Overview sheet", "Executive summary slide"),
    "combined_distribution": ("Summary Charts sheet", "2 distribution slides"),
    "image": ("One sheet per image", "One slide per image"),
    "parameters": ("Methods sheet", "Methods slide"),
    "raw_data": ("Raw sheets — always last", "Appendix slide"),
    "custom_text": ("— (PowerPoint only for now)", "Text slide"),
}


# ======================================================================
# Snapshot + build
# ======================================================================

def _session_ids(state) -> Tuple[str, str]:
    s = state.session
    if s is None:
        return "", ""
    sample = s.meta.sample_id or (s.sample_meta or {}).get("sample_id", "") or ""
    lot = s.meta.lot_number or (s.lot_meta or {}).get("lot_number", "") or ""
    return str(sample), str(lot)


def results_fingerprint(state) -> str:
    """Hash of the numbers a report would be built from (kept grain ids,
    their total area and the calibration per analysed image)."""
    h = hashlib.sha1()
    for im in state.images():
        r = im.result
        if r is None:
            continue
        ids = sorted(int(g.grain_id) for g in r.grains)
        area = round(float(sum(g.area_px for g in r.grains)), 3)
        h.update(json.dumps([im.filename, ids, area, round(float(r.px_per_um or 0.0), 6)])
                 .encode("utf-8"))
    return h.hexdigest()[:16]


def analysed_count(state) -> int:
    return sum(1 for im in state.images() if im.result is not None)


def _filters_text(opts) -> str:
    on = []
    if opts.exclude_border:
        on.append("border grains")
    if opts.exclude_touching_invalid:
        on.append("grains touching excluded regions")
    if opts.exclude_low_contrast:
        on.append("low-contrast / dark false grains")
    if opts.min_area_px:
        on.append(f"area < {opts.min_area_px} px")
    if opts.max_area_px:
        on.append(f"area > {opts.max_area_px} px")
    if opts.max_aspect_ratio:
        on.append(f"aspect ratio > {opts.max_aspect_ratio:g}")
    if opts.min_circularity:
        on.append(f"circularity < {opts.min_circularity:g}")
    return ("Excluded: " + ", ".join(on)) if on else "None"


def session_metadata(state) -> dict:
    s = state.session
    if s is None:
        return {}
    params = {k: v for k, v in (s.params or {}).items() if not isinstance(v, (dict, list))}
    mode = params.pop("detection_mode", "") or s.meta.detector_mode or ""
    params["grain filters"] = _filters_text(s.filters)
    manual = sum(len(im.manual) for im in s.images)
    if manual:
        params["grains removed by hand"] = manual
    return {
        "detection_mode": mode,
        "detection_params": params,
        "instrument": s.meta.instrument or "",
        "magnification": s.meta.magnification or "",
        "session": s.title,
        "project": s.meta.project or (s.project_meta or {}).get("name", ""),
    }


def collect_inputs(state) -> List[ReportImageInput]:
    """GUI-thread snapshot of every analysed image (filtered results)."""
    from ui.workers import snapshot_result
    sample, lot = _session_ids(state)
    out: List[ReportImageInput] = []
    for im in state.images():
        if im.result is None:
            continue
        res = snapshot_result(im.result)
        path = str(im.path) if im.path else im.filename
        out.append(ReportImageInput(image_path=path, result=res, overlay_bgr=res.overlay_image,
                                    sample_id=sample, lot_number=lot))
    return out


def build_model(inputs: Sequence[ReportImageInput], *, title: str, operator: str,
                organization: str = "", logo_path: Optional[str] = None,
                metadata: Optional[dict] = None, asset_dir: Optional[str] = None,
                fingerprint: str = "") -> ReportModel:
    """Pool-thread: ReportModel from snapshots (writes overlay PNGs)."""
    if asset_dir:
        os.makedirs(asset_dir, exist_ok=True)
    meta = dict(metadata or {})
    meta["results_fingerprint"] = fingerprint
    meta.setdefault("exports", [])
    model = ReportModel.from_results(list(inputs), title=title, operator=operator,
                                     organization=organization, logo_path=logo_path,
                                     metadata=meta, asset_dir=asset_dir)
    normalize(model)
    return model


# ======================================================================
# Ordering
# ======================================================================

def image_sections(model: ReportModel) -> List[Section]:
    return [s for s in model.sections if s.type == "image"]


def outline_order(model: ReportModel) -> List[str]:
    """Top-level outline entries as section ids, with the pseudo id
    ``"images"`` standing for the per-image block."""
    secs = sorted(model.sections, key=lambda s: s.order)
    out: List[str] = []
    for s in secs:
        if s.type == "image":
            if "images" not in out:
                out.append("images")
            continue
        out.append(s.id)
    if "images" not in out:
        i = out.index("combined_distribution") + 1 if "combined_distribution" in out else len(out)
        out.insert(i, "images")
    return out


def apply_order(model: ReportModel, top: Sequence[str], image_ids: Sequence[str]) -> None:
    """Write an outline order back into ``Section.order`` / ``ImageSummary.order``.

    Structural rules the renderers rely on are enforced here: Cover first,
    Raw data last (the lab manager's "raw data pages at the end")."""
    top = [t for t in top if t not in ("cover", "raw_data")]
    top = ["cover"] + top + ["raw_data"]
    by_img = {i.id: i for i in model.images}
    ordered_imgs = [by_img[i] for i in image_ids if i in by_img]
    ordered_imgs += [i for i in model.images if i not in ordered_imgs]
    for n, img in enumerate(ordered_imgs, 1):
        img.order = n
    sec_by_id = {s.id: s for s in model.sections}
    img_secs = {s.payload.get("image_id"): s for s in image_sections(model)}
    k = 0
    for t in top:
        if t == "images":
            for img in ordered_imgs:
                s = img_secs.get(img.id)
                if s is not None:
                    s.order = k
                    k += 1
            continue
        s = sec_by_id.get(t)
        if s is not None:
            s.order = k
            k += 1
    model.sections.sort(key=lambda s: s.order)
    model.images.sort(key=lambda i: i.order)


def normalize(model: ReportModel) -> None:
    """Make sure every image has a section and orders are consistent."""
    have = {s.payload.get("image_id") for s in image_sections(model)}
    known = {i.id for i in model.images}
    model.sections = [s for s in model.sections
                      if s.type != "image" or s.payload.get("image_id") in known]
    for img in model.images:
        if img.id not in have:
            model.sections.append(Section(id=f"image_{img.id}", type="image",
                                          title=os.path.basename(img.image_path) or img.id,
                                          order=10_000 + img.order,
                                          payload={"image_id": img.id}))
    for fixed, title in (("parameters", "Methods"), ("raw_data", "Raw Data")):
        if model.get_section(fixed) is None:
            model.sections.append(Section(id=fixed, type=fixed, title=title, order=20_000))
    images_sorted = [i.id for i in sorted(model.images, key=lambda i: i.order)]
    apply_order(model, outline_order(model), images_sorted)


def add_custom_text(model: ReportModel, after: Optional[str] = None,
                    title: str = "Notes", body: str = "") -> Section:
    n = 1
    ids = {s.id for s in model.sections}
    while f"text_{n}" in ids:
        n += 1
    sec = Section(id=f"text_{n}", type="custom_text", title=title, enabled=True,
                  order=0, payload={"body": body})
    model.sections.append(sec)
    top = [t for t in outline_order(model) if t != sec.id]
    anchor = after if after in top else ("parameters" if "parameters" in top else top[-1])
    idx = top.index(anchor) + 1 if anchor != "raw_data" else len(top) - 1
    top.insert(idx, sec.id)
    apply_order(model, top, [i.id for i in sorted(model.images, key=lambda i: i.order)])
    return sec


def remove_section(model: ReportModel, section_id: str) -> None:
    model.sections = [s for s in model.sections if not (s.id == section_id and s.type == "custom_text")]


# ======================================================================
# Refresh-numbers merge
# ======================================================================

def _key(img: ImageSummary) -> str:
    return os.path.basename(img.image_path or img.id).lower()


def merge_refresh(old: ReportModel, new: ReportModel) -> ReportModel:
    """``new`` numbers + every edit made to ``old``.  Image ids of matched
    images are preserved so the page's selection survives a refresh."""
    out = new
    for f in ("title", "operator", "organization", "logo_path", "date", "units", "theme"):
        setattr(out, f, getattr(old, f))
    out.bins = dict(old.bins)
    meta = dict(old.metadata)
    for k in ("detection_mode", "detection_params", "instrument", "magnification", "session",
              "project", "results_fingerprint"):
        if k in new.metadata:
            meta[k] = new.metadata[k]
    out.metadata = meta

    old_by_key = {_key(i): i for i in old.images}
    used_ids = {i.id for i in old.images}
    next_order = max([i.order for i in old.images] or [0]) + 1
    seq = len(used_ids) + 1
    for img in out.images:
        prev = old_by_key.get(_key(img))
        if prev is not None:
            img.id = prev.id
            img.include, img.caption, img.notes, img.order = (prev.include, prev.caption,
                                                              prev.notes, prev.order)
            notes = {int(g["id"]): g["note"] for g in prev.grains if g.get("note")}
            for g in img.grains:
                if int(g["id"]) in notes:
                    g["note"] = notes[int(g["id"])]
        else:
            while f"img_{seq}" in used_ids:
                seq += 1
            img.id = f"img_{seq}"
            used_ids.add(img.id)
            img.order = next_order
            next_order += 1

    old_secs = {s.id: s for s in old.sections if s.type != "image"}
    old_img_secs = {s.payload.get("image_id"): s for s in image_sections(old)}
    secs: List[Section] = []
    for s in out.sections:
        if s.type == "image":
            continue
        prev = old_secs.get(s.id)
        if prev is not None:
            s.enabled, s.title, s.order = prev.enabled, prev.title, prev.order
        secs.append(s)
    for s in old.sections:
        if s.type == "custom_text":
            secs.append(copy.deepcopy(s))
    for img in out.images:
        prev = old_img_secs.get(img.id)
        secs.append(Section(id=f"image_{img.id}", type="image",
                            title=prev.title if prev else os.path.basename(img.image_path),
                            enabled=prev.enabled if prev else True,
                            order=prev.order if prev else 10_000 + img.order,
                            payload={"image_id": img.id}))
    out.sections = secs
    normalize(out)
    return out


# ======================================================================
# Rendering
# ======================================================================

def prepare_render_model(model: ReportModel) -> ReportModel:
    """Copy tuned for the renderers: included images renumbered 1..k in the
    designer's order, grain annotations appended to the image notes."""
    rm = ReportModel.from_dict(copy.deepcopy(model.to_dict()))
    inc = [i for i in sorted(rm.images, key=lambda i: i.order) if i.include]
    exc = [i for i in sorted(rm.images, key=lambda i: i.order) if not i.include]
    for n, img in enumerate(inc + exc, 1):
        img.order = n
    for img in rm.images:
        ann = [f"#{g['id']}: {g['note']}" for g in img.grains if str(g.get("note", "")).strip()]
        if ann:
            img.notes = (img.notes + "\n" if img.notes else "") + "Grain notes — " + "; ".join(ann)
    return rm


def pptx_insert_plan(model: ReportModel) -> List[Tuple[int, Section]]:
    """(slide index, section) for each enabled custom text section, based on
    the slides ``reports.pptx_renderer`` produces for everything before it."""
    included = {i.id for i in model.ordered_images(included_only=True)}
    has_images = bool(included)
    count = 0
    plan: List[Tuple[int, Section]] = []
    for s in sorted(model.sections, key=lambda s: s.order):
        if s.type == "cover":
            count += 1
        elif s.type == "overview_table":
            count += 1
        elif s.type == "combined_distribution":
            count += 2 if (s.enabled and has_images) else 0
        elif s.type == "image":
            count += 1 if s.payload.get("image_id") in included else 0
        elif s.type == "parameters":
            count += 1 if s.enabled else 0
        elif s.type == "custom_text" and s.enabled:
            plan.append((count, s))
            count += 1
    return plan


def add_custom_text_slides(pptx_path: str, model: ReportModel) -> int:
    """Insert the designer's text sections as slides and renumber footers.
    (Interim: belongs in ``reports.pptx_renderer`` once it renders
    ``custom_text``.)"""
    plan = pptx_insert_plan(model)
    if not plan:
        return 0
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from reports import pptx_renderer as pr

    prs = Presentation(pptx_path)
    layout = pr._blank_layout(prs)
    sld_ids = prs.slides._sldIdLst  # noqa: SLF001 — python-pptx has no public reorder API
    for idx, sec in plan:
        slide = prs.slides.add_slide(layout)
        pr._slide_heading(slide, sec.title or "Notes")
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.4))
        tf = box.text_frame
        tf.word_wrap = True
        body = str(sec.payload.get("body", "") or "")
        for n, line in enumerate(body.split("\n") or [""]):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = line
            run.font.size = Pt(16)
            run.font.name = "Calibri"
            run.font.color.rgb = pr.TEXT_DARK
        pr._add_footer(slide, model, 0)
        el = list(sld_ids)[-1]
        sld_ids.remove(el)
        sld_ids.insert(idx, el)
    footer_top = pr.SLIDE_H - Inches(0.4)
    for n, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if shp.has_text_frame and shp.top is not None and shp.top >= footer_top \
                    and shp.text_frame.text.strip().isdigit():
                runs = shp.text_frame.paragraphs[0].runs
                if runs:
                    runs[0].text = str(n)
    prs.save(pptx_path)
    return len(plan)


def render_outputs(model_dict: dict, jobs: Sequence[Tuple[str, str]]) -> List[str]:
    """Pool-thread: render ``[(kind, path), ...]`` (kind = xlsx | pptx)."""
    from reports.excel_renderer import render_excel
    from reports.pptx_renderer import render_pptx
    rm = prepare_render_model(ReportModel.from_dict(model_dict))
    done: List[str] = []
    for kind, path in jobs:
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        if kind == "xlsx":
            render_excel(rm, path)
        else:
            render_pptx(rm, path)
            add_custom_text_slides(path, rm)
        done.append(path)
    return done


def only_image_model(model: ReportModel, image_path: str) -> ReportModel:
    """Copy of ``model`` restricted to one image (Export current image)."""
    m = ReportModel.from_dict(copy.deepcopy(model.to_dict()))
    key = os.path.basename(image_path).lower()
    for img in m.images:
        img.include = _key(img) == key
    return m


# ======================================================================
# Files
# ======================================================================

def slug(text: str, fallback: str = "report") -> str:
    s = re.sub(r"[^\w\-]+", "_", text or "", flags=re.UNICODE).strip("_")
    return (s or fallback)[:60]


def default_export_path(session_path: Path, title: str, kind: str,
                        stamp: Optional[str] = None) -> Path:
    stamp = stamp or datetime.now().strftime("%Y%m%d-%H%M")
    d = Path(session_path) / "exports"
    base = f"{slug(title)}_{stamp}"
    p = d / f"{base}.{kind}"
    n = 2
    while p.exists():
        p = d / f"{base}_{n}.{kind}"
        n += 1
    return p


def copy_logo(src: str, session_path: Path) -> str:
    d = Path(session_path) / REPORT_ASSETS
    d.mkdir(parents=True, exist_ok=True)
    dst = d / ("logo" + Path(src).suffix.lower())
    if Path(src).resolve() != dst.resolve():
        shutil.copyfile(src, dst)
    return str(dst)


def record_export(model: ReportModel, paths: Sequence[str], operator: str) -> None:
    hist = list(model.metadata.get("exports") or [])
    now = datetime.now().strftime("%Y-%m-%d %H:%M")
    for p in paths:
        hist.append({"file": os.path.basename(p), "path": str(p), "time": now,
                     "operator": operator,
                     "kind": "PowerPoint" if str(p).lower().endswith(".pptx") else "Excel"})
    model.metadata["exports"] = hist[-MAX_EXPORT_HISTORY:]


def save_report(session_path: Path, model_dict: dict) -> str:
    """Pool-thread: write report.json through the data layer."""
    from data.session_io import update_session
    update_session(session_path, report=model_dict)
    return datetime.now().strftime("%H:%M")


def load_report(session_path: Path) -> Optional[dict]:
    """Pool-thread: report.json as a dict (None when no report was built)."""
    from data.models import read_json
    p = Path(session_path) / "report.json"
    try:
        d = read_json(p) if p.exists() else {}
    except Exception:
        return None
    return d if d.get("sections") is not None and d.get("images") is not None else None


# ======================================================================
# Preview data (same helpers the Excel renderer uses)
# ======================================================================

def overview_table(model: ReportModel):
    """(header, rows, total_row) exactly as the Overview sheet prints them
    (strings, already formatted with the sheet's number formats)."""
    import numpy as np
    from reports.charts import resolve_units
    from reports.excel_renderer import _combined_unit, _row_size_stats

    images = model.ordered_images(included_only=True)
    au, du = _combined_unit(model, images)
    header = ["#", "Image", "Sample", "Lot", "Grains", f"Mean Area ({au})*",
              f"Median Area ({au})*", f"Std Area ({au})*", f"Mean Diameter ({du})*",
              f"Std Diameter ({du})*", "Units", "Coverage %", "Invalid %", "Mean Circularity",
              "Mean Aspect Ratio", "ASTM G"]
    rows = []
    for n, img in enumerate(sorted(images, key=lambda i: i.order), 1):
        au_i, du_i, mean_a, med_a, std_a, mean_d, std_d = _row_size_stats(model, img)
        rows.append([str(n), os.path.basename(img.image_path), img.sample_id, img.lot_number,
                     str(img.grain_count), f"{mean_a:.2f}", f"{med_a:.2f}", f"{std_a:.2f}",
                     f"{mean_d:.2f}", f"{std_d:.2f}", f"{au_i} / {du_i}",
                     f"{img.grain_coverage_pct:.1f}%", f"{img.invalid_area_pct:.1f}%",
                     f"{img.mean_circularity:.4f}", f"{img.mean_aspect_ratio:.4f}",
                     f"{img.astm_g:.2f}" if img.astm_g is not None else "—"])
    total = None
    grains = [g for img in images for g in img.grains]
    if grains:
        cal = all(i.has_calibration for i in images)
        if cal:
            _, am, _, dm = resolve_units(images[0].px_per_um, model.units)
            areas = np.array([g["area_um2"] for g in grains]) * am
            diams = np.array([g["diameter_um"] for g in grains]) * dm
        else:
            areas = np.array([g["area_px"] for g in grains], dtype=float)
            diams = np.array([g["diameter_px"] for g in grains], dtype=float)
        total = ["", "Combined (all images)", "", "", str(len(grains)),
                 f"{np.mean(areas):.2f}", f"{np.median(areas):.2f}", f"{np.std(areas):.2f}",
                 f"{np.mean(diams):.2f}", f"{np.std(diams):.2f}",
                 f"{au} / {du}" if cal else "px² / px",
                 f"{np.mean([i.grain_coverage_pct for i in images]):.2f}",
                 f"{np.mean([i.invalid_area_pct for i in images]):.2f}",
                 f"{np.mean([g['circularity'] for g in grains]):.4f}",
                 f"{np.mean([g['aspect_ratio'] for g in grains]):.4f}", "—"]
    return header, rows, total


def combined_values(model: ReportModel, kind: str):
    """(values, unit) of all included grains, in report units (as the
    Summary Charts sheet computes them)."""
    from reports.charts import resolve_units
    images = model.ordered_images(included_only=True)
    grains = [g for img in images for g in img.grains]
    if images and all(i.has_calibration for i in images):
        au, am, du, dm = resolve_units(images[0].px_per_um, model.units)
        if kind == "area":
            return [g["area_um2"] * am for g in grains], au
        return [g["diameter_um"] * dm for g in grains], du
    if kind == "area":
        return [g["area_px"] for g in grains], "px²"
    return [g["diameter_px"] for g in grains], "px"


# ======================================================================
# Validation hints
# ======================================================================

def problem_hints(model: ReportModel) -> List[Tuple[str, str, str]]:
    """(severity, problem, fix hint) for ``model.validate()`` findings."""
    out = []
    names = {i.id: os.path.basename(i.image_path) for i in model.images}
    for p in model.validate():
        m = re.match(r"Image (\w+) path does not exist", p)
        if m:
            out.append(("danger", f"{names.get(m.group(1), m.group(1))}: image file is missing",
                        "The file was moved or deleted. Put it back in the session's images "
                        "folder, or untick the image in the outline to leave it out."))
            continue
        m = re.match(r"Image (\w+) overlay path does not exist", p)
        if m:
            out.append(("warning", f"{names.get(m.group(1), m.group(1))}: grain overlay is missing",
                        "Click “Refresh numbers” to regenerate the overlay images."))
            continue
        if p.startswith("title is empty"):
            out.append(("warning", "The report has no title", "Enter a title in Document ▸ Title."))
            continue
        out.append(("warning", p, "Rebuild the report from the session to repair it."))
    if model.logo_path and not os.path.exists(model.logo_path):
        out.append(("warning", "The logo file is missing",
                    "Choose the logo again in Document ▸ Logo."))
    if not model.ordered_images(included_only=True):
        out.append(("warning", "No images are included",
                    "Tick at least one image under “Images” in the outline."))
    return out


__all__ = [
    "REPORT_ASSETS", "SECTION_COLORS", "SECTION_LABELS", "SECTION_TARGETS",
    "results_fingerprint", "analysed_count", "session_metadata", "collect_inputs",
    "build_model", "outline_order", "apply_order", "normalize", "add_custom_text",
    "remove_section", "merge_refresh", "prepare_render_model", "pptx_insert_plan",
    "add_custom_text_slides", "render_outputs", "only_image_model", "slug",
    "default_export_path", "copy_logo", "record_export", "save_report", "load_report",
    "problem_hints", "image_sections",
]
