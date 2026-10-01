"""Batch 4D PowerPoint restructure: slide order, per-part lot charts, grain
density, no counters / footer text / Methods / Appendix / explainer text."""
import os
import re
import sys

import numpy as np
import pytest
from pptx import Presentation
from pptx.util import Emu

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from reports.lot_summary import (
    density_display, grain_density, image_scan_area, lot_summary_data, part_lot_charts,
)
from reports.model import ImageSummary, ReportModel, Section
from reports.pptx_renderer import render_pptx

HIER = [{"key": "project", "label": "Job #", "value": "24-117"},
        {"key": "sample", "label": "Part Number", "value": ""},
        {"key": "lot", "label": "Lot", "value": ""}]
COUNTER = re.compile(r"\(\s*(cont'?d|continued)?\s*\d+\s*/\s*\d+\s*\)", re.I)


def _img(iid, part, lot, n=20, ppu=8.0, order=0, scan_px=None, calibrated=True, cov=None, g=6.0):
    rng = np.random.default_rng(order)
    ds = rng.normal(3, 0.6, n).clip(0.3)
    grains = [{"diameter_um": d, "area_um2": np.pi * (d / 2) ** 2,
               "diameter_px": d * ppu, "area_px": np.pi * (d * ppu / 2) ** 2} for d in ds]
    return ImageSummary(id=iid, image_path="", order=order, px_per_um=ppu if calibrated else 0.0,
                        has_calibration=calibrated, grain_count=n, grains=grains, astm_g=g,
                        valid_area_px=scan_px or 0.0, grain_coverage_pct=cov or 0.0,
                        levels={"sample": part, "lot": lot})


def _model(images, extra=()):
    secs = [Section(id="cover", type="cover", title="Cover", order=0),
            Section(id="ov", type="overview_table", title="Overview", order=1),
            Section(id="ls", type="lot_summary", title="Lot Summary", order=2),
            Section(id="cd", type="combined_distribution", title="Combined", order=3),
            Section(id="params", type="parameters", title="Methods", order=4),
            Section(id="lc", type="lot_comparison", title="Lot comparison", order=5,
                    payload={"parts": [{"part": "P1", "comparison": {"summaries": [], "matrix": []}}]}),
            Section(id="raw", type="raw_data", title="Raw", order=99)] + list(extra)
    return ReportModel(title="Deck", sections=secs, images=images, hierarchy=HIER)


def _deck(parts=3, lots=3, per=2, n=20):
    imgs, k = [], 0
    for p in range(parts):
        for l in range(lots):
            for _ in range(per):
                imgs.append(_img(f"i{k}", f"P{p + 1}", f"P{p + 1}-L{l + 1}", n=n, order=k,
                                 scan_px=1000.0 * 8 * 8))
                k += 1
    return _model(imgs)


def _head(slide):
    return next(r.text for sh in slide.shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs)


def _text(slide):
    return "\n".join(r.text for sh in slide.shapes if sh.has_text_frame
                     for p in sh.text_frame.paragraphs for r in p.runs)


@pytest.fixture(scope="module")
def deck(tmp_path_factory):
    out = str(tmp_path_factory.mktemp("d") / "d.pptx")
    render_pptx(_deck(), out)
    return Presentation(out)


def test_lot_chart_slides_immediately_follow_grain_size_summary(deck):
    heads = [_head(s) for s in deck.slides]
    i = heads.index("Grain Size Summary")
    assert heads[i + 1:i + 7] == ["Mean Grain Diameter by Lot", "Mean Grain Area by Lot",
                                  "Mean Grain Density by Lot",
                                  "Lot Summary — P1", "Lot Summary — P2", "Lot Summary — P3"]
    assert heads[0] == "Deck" and heads[1] == "Contents"
    assert not any(h.startswith("Lot-vs-Lot") for h in heads)      # no duplicate lot-chart slides


def test_lot_chart_slide_has_one_bar_per_lot_and_density_chart(deck):
    s = next(s for s in deck.slides if _head(s) == "Lot Summary — P2")
    charts = [sh.chart for sh in s.shapes if sh.has_chart]
    assert len(charts) == 4
    for c in charts:
        assert [str(x) for x in c.plots[0].categories] == ["P2-L1", "P2-L2", "P2-L3"]
    dens = charts[-1]
    assert dens.chart_title.text_frame.text == "Grain Density by Lot"
    assert dens.value_axis.axis_title.text_frame.text.startswith("Grain Density (grains/")
    assert not any(c.value_axis.axis_title.text_frame.text == "Number of Grains" for c in charts)


def test_all_lots_charts_one_bar_per_lot_with_units(deck):
    data = lot_summary_data(_deck())
    lots = data["lots"]
    expect = {"Mean Grain Diameter by Lot": ("Mean Equivalent Diameter (µm)", "mean_diameter"),
              "Mean Grain Area by Lot": ("Mean Grain Area (µm²)", "mean_area"),
              "Mean Grain Density by Lot": ("Grain Density (grains/", "grain_density")}
    for head, (ytitle, key) in expect.items():
        s = next(s for s in deck.slides if _head(s) == head)
        charts = [sh.chart for sh in s.shapes if sh.has_chart]
        assert len(charts) == 1
        c = charts[0]
        assert c.value_axis.axis_title.text_frame.text.startswith(ytitle)
        assert c.category_axis.axis_title.text_frame.text
        cats = c.plots[0].categories
        assert cats.depth == 2
        assert [str(x) for x in cats.flattened_labels] == [(l["part"], l["lot"]) for l in lots] or \
            [x[-1] for x in cats.flattened_labels] == [l["lot"] for l in lots]
        vals = list(c.plots[0].series[0].values)
        assert len(vals) == 9 and all(v is not None and v > 0 for v in vals)
        if key == "mean_diameter":
            assert vals == pytest.approx([round(l[key], 12) for l in lots], rel=1e-6)
        assert not c.has_legend                       # no legend: the line is not a legend entry
        assert len(c.plots) == 2                      # bars + connecting line
        bars, line = c.plots[0].series[0], c.plots[1].series[0]
        assert list(line.values) == pytest.approx(list(bars.values), abs=1e-3)
        xml = c._chartSpace.xml
        assert '<c:smooth val="0"/>' in xml and '<c:symbol val="circle"/>' in xml   # straight, no fitting


def test_all_lots_charts_continue_without_counter_and_skip_when_disabled(tmp_path):
    imgs = [_img(f"a{k}", "P1", f"L{k}", n=5, order=k) for k in range(25)]
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(imgs), out)
    heads = [_head(s) for s in Presentation(out).slides]
    assert heads.count("Mean Grain Diameter by Lot") == 2
    assert heads.count("Mean Grain Density by Lot") == 2
    m = _model(imgs)
    m.get_section("lot_summary").enabled = False
    render_pptx(m, out)
    assert "Mean Grain Area by Lot" not in [_head(s) for s in Presentation(out).slides]


def test_one_distribution_slide_per_part_with_series_per_lot_and_legend(deck):
    slides = [s for s in deck.slides if _head(s).startswith("Grain Distributions")]
    assert [_head(s) for s in slides] == [f"Grain Distributions — P{k}" for k in (1, 2, 3)]
    for s in slides:
        charts = [sh.chart for sh in s.shapes if sh.has_chart]
        assert len(charts) == 2
        for c in charts:
            assert len(c.plots[0].series) == 3 and c.has_legend


def test_no_slide_heading_has_a_counter(tmp_path):
    # force every continuation path: many parts, long tables, many lots
    imgs = [_img(f"a{k}", "P1", f"L{k % 20}", n=5, order=k) for k in range(60)]
    imgs += [_img(f"b{k}", f"Q{k}", "L1", n=5, order=100 + k) for k in range(20)]
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(imgs), out)
    prs = Presentation(out)
    heads = [_head(s) for s in prs.slides]
    assert len(heads) > 30
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and Emu(sh.top).inches < 1.0:
                assert not COUNTER.search(sh.text_frame.text), sh.text_frame.text
    assert not any(COUNTER.search(h) for h in heads)
    assert heads.count("Contents") >= 1


def test_footer_has_only_the_page_number(deck):
    h = Emu(deck.slide_height).inches
    for n, s in enumerate(deck.slides, start=1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame
                 and sh.text_frame.text.strip() and abs(Emu(sh.top).inches - (h - 0.32)) < 0.05]
        assert texts == [str(n)]
        assert "parts" not in "".join(texts) and "24-117" not in "".join(texts)


def test_removed_slides_and_explainer_text_are_gone_even_with_old_sections(deck):
    # the fixture model still carries parameters / lot_comparison / raw_data sections
    heads = [_head(s) for s in deck.slides]
    assert not any(h.startswith(("Methods", "Appendix", "Lot Comparison")) for h in heads)
    joined = "\n".join(_text(s) for s in deck.slides)
    assert "smaller than this size" not in joined
    assert "the median" not in joined and "90 % are smaller" not in joined
    assert "Detection mode" not in joined
    two_table = [s for s in deck.slides if sum(1 for sh in s.shapes if sh.has_table) >= 2]
    assert not two_table


def test_contents_lists_new_structure_with_working_links(deck):
    ids = [s.slide_id for s in deck.slides]
    rows = {}
    for s in deck.slides:
        if _head(s) != "Contents":
            continue
        for sh in s.shapes:
            if sh.has_text_frame and sh.click_action.target_slide is not None:
                m = re.match(r"^Pages? (\d+)(?:–(\d+))?\s+·\s+(.+)$", sh.text_frame.text)
                rows[m.group(3)] = (int(m.group(1)), int(m.group(2) or m.group(1)),
                                    ids.index(sh.click_action.target_slide.slide_id))
    assert rows["Lot summary, all lots"][0] == rows["Grain size summary"][1] + 1
    assert rows["Lot summary, all lots"][1] == rows["Lot summary, all lots"][0] + 2
    assert rows["Lot summary by part"][0] == rows["Lot summary, all lots"][1] + 1
    assert "Lot summary by part" in rows and "Grain distributions by part" in rows
    assert not any(k in rows for k in ("Appendix", "Methods & parameters", "Lot comparison by part"))
    for a, b, tgt in rows.values():
        assert tgt == a - 1 and 1 <= a <= b <= len(ids)


# ---------------------------------------------------------------------------
# Grain density (Qt-free helper)
# ---------------------------------------------------------------------------

def test_density_is_grains_per_scan_area_known_values():
    # 50 grains in a 100 x 100 um scan (8 px/um -> 800 x 800 px)
    im = _img("a", "P", "L", n=50, scan_px=800 * 800)
    assert image_scan_area(im, True) == pytest.approx(10000.0)
    assert grain_density([im], True) == pytest.approx(50 / 10000.0)
    # a lot = total grains / total scan area (not the mean of per-image densities)
    im2 = _img("b", "P", "L", n=10, order=1, scan_px=800 * 800 * 4)
    assert grain_density([im, im2], True) == pytest.approx(60 / 50000.0)


def test_scan_area_falls_back_to_full_frame_value_and_coverage():
    # detector stores the whole frame as the valid area when no scan rect is set
    full = _img("a", "P", "L", n=40, scan_px=1024 * 768)
    assert image_scan_area(full, False) == pytest.approx(1024 * 768)
    # older saved reports (no valid_area_px): derive from coverage = grain area / scan area
    old = _img("o", "P", "L", n=40, scan_px=None)
    old.grain_coverage_pct = 80.0
    total_px = sum(g["area_px"] for g in old.grains)
    assert image_scan_area(old, False) == pytest.approx(total_px / 0.8)
    old.valid_area_um2 = 500.0
    assert image_scan_area(old, True) == pytest.approx(500.0)
    # nothing known -> None (never a made-up number)
    assert image_scan_area(_img("n", "P", "L", scan_px=None), False) is None
    assert grain_density([_img("n", "P", "L", scan_px=None)], False) is None


def test_density_skips_images_without_scan_area_in_numerator_and_denominator():
    ok = _img("a", "P", "L", n=20, scan_px=64 * 1000)           # 1000 um^2
    unknown = _img("b", "P", "L", n=500, order=1, scan_px=None)
    assert grain_density([ok, unknown], True) == pytest.approx(20 / 1000.0)


def test_density_display_unit_chosen_by_magnitude_and_honest_when_uncalibrated():
    vals, unit, fmt = density_display([0.05, 0.02, None], True)
    assert unit == "grains/µm²" and vals[:2] == [0.05, 0.02] and vals[2] is None
    vals, unit, fmt = density_display([2e-4, 1e-4], True)
    assert unit == "grains/mm²" and vals == pytest.approx([200.0, 100.0])
    vals, unit, fmt = density_display([1e-3], False)
    assert unit.startswith("grains/Mpx") and "uncalibrated" in unit and vals == pytest.approx([1000.0])


def test_lot_summary_density_uses_lot_scan_area_and_chart_axis_title():
    imgs = [_img("a", "P1", "L1", n=30, order=1, scan_px=64 * 2000),
            _img("b", "P1", "L1", n=10, order=2, scan_px=64 * 2000),
            _img("c", "P1", "L2", n=5, order=3, scan_px=64 * 1000)]
    data = lot_summary_data(_model(imgs), imgs)
    by_lot = {s["lot"]: s for s in data["lots"]}
    assert by_lot["L1"]["grain_density"] == pytest.approx(40 / 4000.0)
    assert by_lot["L2"]["grain_density"] == pytest.approx(5 / 1000.0)
    dens = part_lot_charts(data, "P1")[-1]
    assert dens["id"] == "grain_density"
    assert dens["y_title"] == "Grain Density (grains/µm²)"
    assert dens["values"] == pytest.approx([0.01, 0.005])


def test_excluded_images_do_not_count_toward_density():
    a = _img("a", "P1", "L1", n=30, order=1, scan_px=64 * 1000)
    b = _img("b", "P1", "L1", n=30, order=2, scan_px=64 * 1000)
    b.include = False
    m = _model([a, b])
    data = lot_summary_data(m)                   # default: included images only
    assert data["lots"][0]["grain_density"] == pytest.approx(30 / 1000.0)


def test_uncalibrated_density_axis_is_labelled_per_megapixel():
    imgs = [_img("a", "P1", "L1", n=30, order=1, scan_px=1000 * 1000, calibrated=False),
            _img("b", "P1", "L2", n=60, order=2, scan_px=1000 * 1000, calibrated=False)]
    data = lot_summary_data(_model(imgs), imgs)
    dens = part_lot_charts(data, "P1")[-1]
    assert dens["y_title"] == "Grain Density (grains/Mpx, uncalibrated)"
    assert dens["values"] == pytest.approx([30.0, 60.0])


def test_old_report_json_without_valid_area_px_still_loads():
    m = _deck(parts=1, lots=1, per=1)
    d = m.to_dict() if hasattr(m, "to_dict") else None
    import json
    raw = json.loads(m.to_json())
    for im in raw["images"]:
        im.pop("valid_area_px", None)
    m2 = ReportModel.from_json(json.dumps(raw))
    assert m2.images[0].valid_area_px == 0.0
