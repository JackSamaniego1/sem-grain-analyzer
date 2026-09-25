"""
UX-09 / UX-13 wave 2: multi-lot report from the Analyze page.

Full journey: load a part (3 lots) into the analyzer -> auto-find -> analyse
-> edit (delete a grain) -> "Export report…" (designer + direct XLSX/PPTX)
-> reopen, with the lot comparison section in the designer outline and
preview, overlay opacity + chart defaults honoured, and a selection scope.
"""
from pathlib import Path

import pytest

pytest.importorskip("pytestqt")

from data.models import AppSettings  # noqa: E402
from data.settings import save_settings  # noqa: E402
from tests.test_ui_ux_v301 import TIMEOUT, _lots, _shell  # noqa: E402
from tests.ui_shell_helpers import confirm_setup  # noqa: E402


@pytest.fixture
def env(tmp_path, monkeypatch, qapp):
    from ui.design.theme import apply_theme, set_reduced_motion
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    root = tmp_path / "ws"
    save_settings(AppSettings(workspace_root=str(root), operator="Tester", theme="dark"))
    set_reduced_motion(True)
    apply_theme(qapp, "dark")
    yield root
    set_reduced_motion(False)


def _analysed_part(env, qtbot, n_lots=3, per_lot=3):
    sample = _lots(env, n_lots, per_lot, sample="7718-A")
    shell = _shell(qtbot)
    st = shell.state
    shell.load_into_analyzer([sample])
    qtbot.waitUntil(lambda: st.session is not None and st.session.multi
                    and not st.is_loading(), timeout=TIMEOUT)
    shell.analyze.params.set_mode("threshold")
    confirm_setup(shell, qtbot)
    with qtbot.waitSignal(shell.analyze.queue.queue_finished, timeout=TIMEOUT):
        shell.analyze.btn_all.click()
    qtbot.waitUntil(lambda: all(im.result is not None for im in st.images())
                    and not st.is_filtering(), timeout=TIMEOUT)
    return shell, sample


def _wait_model(shell, qtbot):
    from ui.pages import report_builder as rb
    r = shell.reports
    qtbot.waitUntil(lambda: r.model is not None and rb.is_multi_model(r.model)
                    and not r.is_busy(), timeout=TIMEOUT)
    return r.model


def test_multi_lot_report_journey(env, qtbot, tmp_path):
    from openpyxl import load_workbook
    from ui.pages import report_builder as rb
    from ui.pages.report_preview import LotComparisonPreview
    shell, sample = _analysed_part(env, qtbot)
    st, a = shell.state, shell.analyze
    # baseline lot per material (Projects > Compare lots star)
    lot1 = Path(sample) / "L-1"
    assert (lot1 / "lot.json").exists()
    st.workspace.set_baseline_lot(lot1, True)
    # edit: delete one grain by hand (undoable) before reporting
    im = st.images()[0]
    st.set_current_image(im.uid)
    gid = int(im.result.grains[0].grain_id)
    n0 = im.result.grain_count
    assert st.delete_grains(im.uid, [gid])
    qtbot.waitUntil(lambda: not st.is_filtering() and im.result.grain_count == n0 - 1,
                    timeout=TIMEOUT)
    st.set_overlay_opacity(0.55)

    # --- results table: the Export report button + menu
    a.show_table_view()
    t = a.table
    t.rebuild()
    assert t.report_btn.isVisible() and t.report_btn.isEnabled()
    menu = t.show_report_menu()
    texts = [x.text() for x in menu.actions() if x.text()]
    assert texts[0].startswith("Everything loaded: 3 lots · 9 analysed images")
    assert "Open in report designer…" in texts and "Export Excel + PowerPoint" in texts
    shell.resize(1500, 950)
    qtbot.wait(50)
    shell.grab().save(str(_shot(tmp_path, "wave2_1_results_table.png")))
    menu.popup(t.report_btn.mapToGlobal(t.report_btn.rect().bottomLeft()))
    qtbot.wait(100)
    menu.grab().save(str(_shot(tmp_path, "wave2_1b_export_menu.png")))
    menu.hide()

    # --- open in the designer
    with qtbot.waitSignal(t.report_requested, timeout=2000):
        next(x for x in menu.actions() if x.text() == "Open in report designer…").trigger()
    model = _wait_model(shell, qtbot)
    assert shell.current_page() == "reports"
    assert len(model.images) == 9
    assert abs(model.overlay_opacity - 0.55) < 1e-6
    assert model.chart_options == rb.chart_options_arg(st)
    assert {i.levels.get("lot") for i in model.images} == {"L-1", "L-2", "L-3"}
    sec = model.get_section("lot_comparison")
    assert sec is not None and sec.enabled
    part = sec.payload["parts"][0]
    assert part["lots"] == ["L-1", "L-2", "L-3"] and part["baseline"] == "L-1"
    assert all(img.overlay_path and Path(img.overlay_path).exists() for img in model.images)
    # the edit is in the report
    g0 = next(i for i in model.images if Path(i.image_path).name == im.filename
              and i.levels.get("lot") == "L-1")
    assert g0.grain_count == n0 - 1
    # outline + live preview of the lot comparison section
    r = shell.reports
    it = r.outline.item_for(("section", "lot_comparison"))
    assert it is not None and it.text(0) == "Lot comparison"
    r.select(("section", "lot_comparison"))
    qtbot.waitUntil(lambda: isinstance(r.current_preview(), LotComparisonPreview), timeout=5000)
    pv = r.current_preview()
    assert len(pv.cards) == 1
    verdicts = pv.verdict_texts()
    assert len(verdicts) == 2 and all(v in ("Equivalent", "Not equivalent", "Inconclusive")
                                      for v in verdicts)
    qtbot.wait(100)
    shell.grab().save(str(_shot(tmp_path, "wave2_2_designer_lot_comparison.png")))

    # --- direct export of both files from the Analyze page
    shell.go("analyze")
    with qtbot.waitSignal(r.exported, timeout=TIMEOUT) as blk:
        shell.report_from_analyzer(None, "both")
    paths = blk.args[0]
    assert sorted(Path(p).suffix for p in paths) == [".pptx", ".xlsx"]
    assert all(Path(p).exists() and Path(p).stat().st_size > 0 for p in paths)
    qtbot.waitUntil(lambda: not r.is_busy(), timeout=TIMEOUT)
    qtbot.wait(400)
    shell.grab().save(str(_shot(tmp_path, "wave2_3_exported_toast.png")))
    wb = load_workbook(next(p for p in paths if p.endswith(".xlsx")), read_only=True)
    assert any("Lot Comparison" in n for n in wb.sheetnames), wb.sheetnames
    wb.close()

    # --- scope: only the selected lot group
    t.set_group_level("lot")
    t.rebuild()
    grp = next(g for g in t.group_items() if g.child(0).text(3) == "L-2")
    t.tree.clearSelection()
    grp.setSelected(True)
    scope = t.selected_scope()
    assert scope is not None and len(scope) == 1
    shell.report_from_analyzer(scope, "designer")
    qtbot.waitUntil(lambda: r.model is not None and rb.model_scope(r.model) == scope
                    and not r.is_busy(), timeout=TIMEOUT)
    assert {i.levels.get("lot") for i in r.model.images} == {"L-2"}
    # back to everything; the report is saved with the first lot
    shell.report_from_analyzer(None, "designer")
    qtbot.waitUntil(lambda: r.model is not None and rb.model_scope(r.model) is None
                    and len(r.model.images) == 9 and not r.is_busy(), timeout=TIMEOUT)
    r.set_section_enabled("combined_distribution", True)
    r.flush()
    st.flush()
    shell.close()

    # --- reopen: the multi-lot report comes back with the job
    shell2 = _shell(qtbot)
    shell2.load_into_analyzer([sample])
    st2 = shell2.state
    qtbot.waitUntil(lambda: st2.session is not None and not st2.is_loading(), timeout=TIMEOUT)
    assert all(im.result is not None for im in st2.images())
    shell2.go("reports")
    m2 = _wait_model(shell2, qtbot)
    assert m2.get_section("lot_comparison") is not None and len(m2.images) == 9
    assert shell2.reports._fingerprint() == m2.metadata["results_fingerprint"]
    shell2.close()


def _shot(tmp_path, name: str) -> Path:
    import os
    d = os.environ.get("GRAIN_SHOTS_DIR")
    return Path(d) / name if d else tmp_path / name
