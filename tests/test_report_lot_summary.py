"""Lot Summary (UPDATE 4 item 15): Qt-free data API, Excel sheet, PPTX slides."""
import os
import re
import sys
import zipfile

import numpy as np
import openpyxl
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from reports.excel_renderer import render_excel
from reports.lot_summary import lot_summary_data, id_cells, ASTM_AVG_FOOTNOTE
from reports.model import ImageSummary, ReportModel, Section
from reports.pptx_renderer import (
    LOT_CHART_TITLE, LOT_SUMMARY_TITLE, MAX_DATA_ROWS, MAX_LOTS_PER_CHART, render_pptx,
)

HIER = [{"key": "sample", "label": "Part Number", "value": ""},
        {"key": "lot", "label": "Lot", "value": ""}]
NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
ROW = re.compile(r"^(Page|Pages) (\d+)(?:–(\d+))?\s+·\s+(.+)$")


def _img(iid, part, lot, diams_um, ppu=8.0, order=0, calibrated=True, g=None):
    grains = [{"diameter_um": d, "area_um2": np.pi * (d / 2) ** 2,
               "diameter_px": d * ppu, "area_px": np.pi * (d * ppu / 2) ** 2,
               "circularity": 0.8, "aspect_ratio": 1.2} for d in diams_um]
    return ImageSummary(id=iid, image_path="", order=order, px_per_um=ppu, has_calibration=calibrated,
                        grain_count=len(grains), grains=grains, astm_g=g,
                        levels={"sample": part, "lot": lot})


def _model(images, units="auto", summary=True, extra=()):
    secs = [Section(id="cover", type="cover", title="Cover", order=0),
            Section(id="ov", type="overview_table", title="Overview", order=1)]
    if summary:
        secs.append(Section(id="lot_summary", type="lot_summary", title="Lot Summary", order=2))
    secs += list(extra)
    return ReportModel(title="T", sections=secs, images=images, hierarchy=HIER, units=units)


def _job(parts=2, lots=3, per=2, seed=0):
    rng = np.random.default_rng(seed)
    imgs, k = [], 0
    for p in range(parts):
        for l in range(lots):
            for _ in range(per):
                d = rng.normal(2 + 0.5 * l + p, 0.4, 30).clip(0.2)
                imgs.append(_img(f"i{k}", f"P{p + 1}", f"L{l + 1}", d, order=k, g=5.0 + l))
                k += 1
    return imgs


def _lots(n, per=1, **kw):
    rng = np.random.default_rng(1)
    return [_img(f"i{k}_{j}", "P1", f"Lot-{k + 1}", rng.normal(2 + 0.1 * k, 0.4, 25).clip(0.2),
                 order=k * per + j, g=6.0, **kw) for k in range(n) for j in range(per)]


def _text(slide):
    return "\n".join(r.text for sh in slide.shapes if sh.has_text_frame
                     for p in sh.text_frame.paragraphs for r in p.runs)


def _head(slide):
    return next(r.text for sh in slide.shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs)


def _pptx(tmp_path, model):
    out = str(tmp_path / "d.pptx")
    render_pptx(model, out)
    return Presentation(out)


def _xlsx_chart_xml(path):
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if n.startswith("xl/charts/chart"))
        return [z.read(n).decode("utf8") for n in names]


# ---------------------------------------------------------------------------
# Data API
# ---------------------------------------------------------------------------

def test_data_shape_and_totals():
    m = _model(_job(parts=2, lots=3, per=2))
    d = lot_summary_data(m)
    assert d["has_lots"] and d["multi_part"]
    assert len(d["lots"]) == 6 and len(d["parts"]) == 2
    assert d["total"]["n_images"] == 12 and d["total"]["n_grains"] == 12 * 30
    kinds = [r["kind"] for r in d["rows"]]
    assert kinds == ["lot", "lot", "lot", "part", "lot", "lot", "lot", "part", "total"]
    assert id_cells(d["rows"][-1]) == ("JOB TOTAL", "")
    assert [c["id"] for c in d["charts"]] == ["mean_diameter", "median_diameter", "mean_area",
                                              "n_grains", "astm_g"]
    lot0 = d["lots"][0]
    grains = [g["diameter_um"] for i in m.images[:2] for g in i.grains]
    assert abs(lot0["mean_diameter"] - np.mean(grains)) < 1e-9
    assert abs(lot0["median_diameter"] - np.median(grains)) < 1e-9
    assert lot0["d10"] < lot0["median_diameter"] < lot0["d90"]


def test_data_units_nm_and_uncalibrated_px():
    nm = lot_summary_data(_model(_lots(2), units="nm"))
    assert nm["units"]["length"] == "nm" and "nm" in nm["units"]["area"]
    assert "(nm)" in nm["charts"][0]["y_title"]
    um = lot_summary_data(_model(_lots(2), units="um"))
    assert abs(nm["lots"][0]["mean_diameter"] - um["lots"][0]["mean_diameter"] * 1000) < 1e-6
    px = lot_summary_data(_model(_lots(2, calibrated=False)))
    assert px["units"] == {"calibrated": False, "length": "px", "area": "px²"}
    assert "(px)" in px["charts"][0]["y_title"]


def test_trend_rules_one_two_many_lots():
    assert all(c["trend"] is None for c in lot_summary_data(_model(_lots(1)))["charts"])
    two = lot_summary_data(_model(_lots(2)))["charts"][0]
    assert two["trend"] == two["values"]
    many = lot_summary_data(_model(_lots(20)))["charts"][0]
    assert many["trend"] == many["values"]          # joins the lot values, no fit


def test_line_joins_lot_values_across_all_parts_in_order():
    imgs = _job(parts=1, lots=3, per=1) + [_img("solo", "P9", "Only", [2.0, 3.0, 2.5], order=99)]
    for ch in lot_summary_data(_model(imgs))["charts"]:
        assert ch["trend"] == ch["values"] and len(ch["trend"]) == 4


def test_zero_grain_lot_and_missing_astm():
    imgs = _lots(3)
    imgs[1] = _img("empty", "P1", "Lot-2", [], order=1, g=None)
    d = lot_summary_data(_model(imgs))
    empty = d["lots"][1]
    assert empty["n_grains"] == 0 and empty["mean_diameter"] is None and empty["d10"] is None
    ch = {c["id"]: c for c in d["charts"]}
    assert ch["mean_diameter"]["values"][1] is None and ch["n_grains"]["values"][1] == 0
    # line joins the lot values; the empty lot is a gap
    assert ch["mean_diameter"]["trend"] == ch["mean_diameter"]["values"]
    assert ch["mean_diameter"]["trend"][1] is None
    for i in imgs:
        i.astm_g = None
    assert "astm_g" not in [c["id"] for c in lot_summary_data(_model(imgs))["charts"]]


def test_no_lot_values_means_has_lots_false():
    imgs = [_img("a", "P1", "", [1.0, 2.0])]
    assert lot_summary_data(_model(imgs))["has_lots"] is False


def test_long_lot_names_shortened_for_axis_full_in_table():
    long = "LOT-" + "X" * 80
    imgs = [_img("a", "P1", long, [1.0, 2.0]), _img("b", "P1", "L2", [1.5, 2.5], order=1)]
    d = lot_summary_data(_model(imgs))
    assert d["lots"][0]["lot"] == long
    assert len(d["charts"][0]["categories"][0]["label"]) <= 24


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def _sheet(tmp_path, model):
    out = str(tmp_path / "r.xlsx")
    render_excel(model, out)
    return out, openpyxl.load_workbook(out)["Lot Summary"]


def test_excel_job_summary_table_rows_and_total(tmp_path):
    m = _model(_job(parts=2, lots=3, per=2))
    out, ws = _sheet(tmp_path, m)
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5, max_col=10)]
    body = []
    for r in rows:
        if r[0] is None and r[1] is None:
            break
        body.append(r)
    assert body[-1][0] == ASTM_AVG_FOOTNOTE     # footnote sits right under the table
    body = body[:-1]
    assert len(body) == 9                     # 6 lots + 2 part totals + JOB TOTAL
    assert body[-1][0] == "JOB TOTAL" and body[-1][2] == 12 and body[-1][3] == 360
    assert body[3][1] == "All lots" and body[3][0] == "P1"
    header = [c.value for c in ws[4]][:10]
    assert "Mean Diameter (µm)" in header[4] and header[5].startswith("Median Diameter")


def test_excel_lot_charts_are_combo_with_units_in_axes(tmp_path):
    m = _model(_job(parts=1, lots=4, per=1))
    out, _ = _sheet(tmp_path, m)
    xml = [x for x in _xlsx_chart_xml(out) if "by Lot" in x]
    assert len(xml) == 5
    for x in xml:
        assert "<c:barChart>" in x and "<c:lineChart>" in x
        assert "Lot values" in x and "linear fit" not in x
    joined = "".join(xml)
    assert "Mean Equivalent Diameter (µm)" in joined
    assert "Mean Grain Area (µm²)" in joined
    assert "Number of Grains" in joined and "ASTM G Number" in joined
    assert "<c:multiLvlStrRef>" not in joined           # one part -> flat labels


def test_excel_single_lot_has_no_trend_line_and_no_crash(tmp_path):
    out, ws = _sheet(tmp_path, _model(_lots(1)))
    xml = [x for x in _xlsx_chart_xml(out) if "by Lot" in x]
    assert len(xml) == 5 and all("<c:lineChart>" not in x for x in xml)


def test_excel_two_parts_use_two_level_categories(tmp_path):
    out, _ = _sheet(tmp_path, _model(_job(parts=2, lots=2, per=1)))
    xml = "".join(x for x in _xlsx_chart_xml(out) if "by Lot" in x)
    assert "<c:multiLvlStrRef>" in xml
    assert "Part Number / Lot" in xml


def test_excel_many_lots_rotate_labels_and_widen(tmp_path):
    out, ws = _sheet(tmp_path, _model(_lots(20)))
    xml = [x for x in _xlsx_chart_xml(out) if "by Lot" in x]
    assert len(xml) == 5 and all('rot="-2700000"' in x for x in xml)
    anchors = [c.anchor._from.col for c in ws._charts if c.anchor is not None]
    assert 6 not in anchors                     # one full-width chart per row


def test_excel_zero_grain_lot_px_and_long_names_render(tmp_path):
    imgs = _lots(3, calibrated=False)
    imgs[1] = _img("e", "P1", "Lot-2", [], order=1, calibrated=False)
    imgs[2].levels["lot"] = "L" * 90
    out, ws = _sheet(tmp_path, _model(imgs))
    hdr = [c.value for c in ws[4]][:10]
    assert "(px)" in hdr[4]
    xml = "".join(_xlsx_chart_xml(out))
    assert "Mean Equivalent Diameter (px)" in xml
    assert any(r[3].value == 0 for r in ws.iter_rows(min_row=5, max_row=8))


def test_excel_nm_units_respected(tmp_path):
    out, ws = _sheet(tmp_path, _model(_lots(2), units="nm"))
    assert "(nm)" in [c.value for c in ws[4]][4]
    assert "Mean Grain Area (nm²)" in "".join(_xlsx_chart_xml(out))


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------

def _slides(prs, prefix):
    return [s for s in prs.slides if _head(s).startswith(prefix)]


def _chart_xml(chart):
    return chart._chartSpace


def test_pptx_lot_summary_table_and_charts_present(tmp_path):
    prs = _pptx(tmp_path, _model(_job(parts=2, lots=3, per=1)))
    tables = _slides(prs, LOT_SUMMARY_TITLE)
    assert len(tables) == 1
    tbl = next(sh.table for sh in tables[0].shapes if sh.has_table)
    assert len(tbl.rows) == 1 + 9
    assert tbl.cell(9, 0).text == "JOB TOTAL"
    assert "Mean Diameter (µm)" in tbl.cell(0, 4).text
    charts_slides = _slides(prs, LOT_CHART_TITLE)
    assert [_head(s) for s in charts_slides] == ["Lot Summary — P1", "Lot Summary — P2"]
    for s in charts_slides:                    # one slide per part, ONE BAR PER LOT
        charts = [sh.chart for sh in s.shapes if sh.has_chart]
        assert len(charts) == 4               # ASTM G, mean diameter, mean area, grain density
        titles = [c.chart_title.text_frame.text for c in charts]
        assert titles == ["ASTM Grain Size Number by Lot", "Mean Grain Diameter by Lot",
                          "Mean Grain Area by Lot", "Grain Density by Lot"]
        for c in charts:
            assert [str(x) for x in c.plots[0].categories] == ["L1", "L2", "L3"]
            assert c.category_axis.axis_title.text_frame.text == "Lot"
            assert not c._chartSpace.findall(".//c:lineChart", NS)
        assert charts[1].value_axis.axis_title.text_frame.text == "Mean Equivalent Diameter (µm)"
        assert charts[2].value_axis.axis_title.text_frame.text == "Mean Grain Area (µm²)"
        assert charts[3].value_axis.axis_title.text_frame.text.startswith("Grain Density (grains/")
        assert "Number of Grains" not in [c.value_axis.axis_title.text_frame.text for c in charts]


def test_astm_average_footnote_in_data_and_pptx(tmp_path):
    d = lot_summary_data(_model(_job(parts=2, lots=2, per=1)))
    assert d["footnotes"] == [ASTM_AVG_FOOTNOTE]
    assert "average of the image values" in d["footnotes"][0]
    # no subtotal/total ASTM value -> no footnote
    assert lot_summary_data(_model([_img("a", "P1", "L1", [1.0, 2.0])]))["footnotes"] == []
    prs = _pptx(tmp_path, _model(_lots(20)))
    t1, t2 = _slides(prs, LOT_SUMMARY_TITLE)
    assert ASTM_AVG_FOOTNOTE not in _text(t1)   # page 1 has lot rows only
    assert ASTM_AVG_FOOTNOTE in _text(t2)       # page with JOB TOTAL


def test_pptx_single_lot_no_trend(tmp_path):
    prs = _pptx(tmp_path, _model(_lots(1)))
    charts = [sh.chart for s in _slides(prs, LOT_CHART_TITLE) for sh in s.shapes if sh.has_chart]
    assert len(charts) == 4
    assert all(not c._chartSpace.findall(".//c:lineChart", NS) for c in charts)
    assert all(not c.has_legend for c in charts)


def test_pptx_two_lots_one_bar_per_lot_no_legend(tmp_path):
    prs = _pptx(tmp_path, _model(_lots(2)))
    slides = _slides(prs, LOT_CHART_TITLE)
    assert len(slides) == 1
    charts = [sh.chart for sh in slides[0].shapes if sh.has_chart]
    assert all(len(list(c.plots[0].categories)) == 2 for c in charts)
    assert all(not c.has_legend for c in charts)


def test_pptx_table_paginates_at_14_rows_with_total_last(tmp_path):
    prs = _pptx(tmp_path, _model(_lots(20)))
    tables = _slides(prs, LOT_SUMMARY_TITLE)
    assert len(tables) == 2                     # 20 lots + JOB TOTAL = 21 rows -> 14 + 7
    assert _head(tables[0]) == _head(tables[1]) == LOT_SUMMARY_TITLE     # no "(i/N)" counters
    t1 = next(sh.table for sh in tables[0].shapes if sh.has_table)
    t2 = next(sh.table for sh in tables[1].shapes if sh.has_table)
    assert len(t1.rows) - 1 == MAX_DATA_ROWS
    assert t2.cell(len(t2.rows) - 1, 0).text == "JOB TOTAL"


def test_pptx_many_lots_tilt_labels_and_paginate_charts(tmp_path):
    n = MAX_LOTS_PER_CHART + 6
    prs = _pptx(tmp_path, _model(_lots(n)))
    slides = _slides(prs, LOT_CHART_TITLE)
    assert len(slides) == 2                     # minimum split of one part's 18 lots (9 + 9)
    charts = [sh.chart for sh in slides[0].shapes if sh.has_chart]
    body = charts[0]._chartSpace.find(".//c:catAx/c:txPr/{*}bodyPr", NS)
    assert body is not None and body.get("rot") == "-2700000"
    assert len(charts[0].plots[0].categories) == 9
    assert _head(slides[0]) == _head(slides[1]) == "Lot Summary — P1"   # plain repeated title


def test_pptx_zero_grain_lot_and_missing_calibration(tmp_path):
    imgs = _lots(3, calibrated=False)
    imgs[1] = _img("e", "P1", "Lot-2", [], order=1, calibrated=False)
    prs = _pptx(tmp_path, _model(imgs))
    tbl = next(sh.table for s in _slides(prs, LOT_SUMMARY_TITLE) for sh in s.shapes if sh.has_table)
    assert "(px)" in tbl.cell(0, 4).text
    assert tbl.cell(2, 3).text == "0" and tbl.cell(2, 4).text == "–"
    charts = [sh.chart for s in _slides(prs, LOT_CHART_TITLE) for sh in s.shapes if sh.has_chart]
    assert charts[1].value_axis.axis_title.text_frame.text == "Mean Equivalent Diameter (px)"


def test_pptx_long_lot_names_truncated_in_table_with_notes(tmp_path):
    imgs = _lots(2)
    imgs[0].levels["lot"] = "LOT-" + "Y" * 80
    prs = _pptx(tmp_path, _model(imgs))
    s = _slides(prs, LOT_SUMMARY_TITLE)[0]
    tbl = next(sh.table for sh in s.shapes if sh.has_table)
    assert tbl.cell(1, 1).text.endswith("…") and len(tbl.cell(1, 1).text) < 30
    assert "Y" * 80 in s.notes_slide.notes_text_frame.text


def test_pptx_disabled_or_lotless_lot_summary_adds_no_slides(tmp_path):
    prs = _pptx(tmp_path, _model(_lots(2), summary=False))
    assert not _slides(prs, LOT_SUMMARY_TITLE) and not _slides(prs, LOT_CHART_TITLE)
    lotless = [_img("a", "P1", "", [1.0, 2.0])]
    prs = _pptx(tmp_path, _model(lotless))
    assert not _slides(prs, LOT_SUMMARY_TITLE)


def test_pptx_contents_page_lists_lot_summary_with_correct_pages(tmp_path):
    prs = _pptx(tmp_path, _model(_job(parts=2, lots=3, per=1)))
    slides = list(prs.slides)
    ids = [s.slide_id for s in slides]
    found = {}
    for s in slides:
        if not _head(s).startswith("Contents"):
            continue
        for sh in s.shapes:
            if sh.has_text_frame and sh.click_action.target_slide is not None:
                m = ROW.match(sh.text_frame.text)
                found[m.group(4)] = (int(m.group(2)), int(m.group(3) or m.group(2)),
                                     ids.index(sh.click_action.target_slide.slide_id))
    a, b, tgt = found["Job summary by part and lot"]
    assert _head(slides[a - 1]).startswith(LOT_SUMMARY_TITLE) and tgt == a - 1 and a == b
    a, b, tgt = found["Lot summary by part"]
    assert _head(slides[a - 1]).startswith(LOT_CHART_TITLE) and tgt == a - 1
    assert _head(slides[b - 1]).startswith(LOT_CHART_TITLE) and b - a == 1   # one slide per part
    # footers carry the final page numbers
    last = [sh.text_frame.text for sh in slides[a - 1].shapes if sh.has_text_frame][-1]
    assert int(last) == a
