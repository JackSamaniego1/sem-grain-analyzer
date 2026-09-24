"""INN-27: lot statistics, ASTM E112 95 % CI / %RA, field include/exclude,
lot field collection, and the report lot block / PPTX tile."""
import math
import os
import sys

import numpy as np
import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from core.grain_detector import AnalysisResult, GrainResult
from core.metrics import (FieldResult, SampleStatistics, StatsConfig, fields_needed,
                          pooled_ecds_um, sample_statistics)
from data.catalog import Catalog, fields_for_lot
from data.models import AppSettings, ImageManifestEntry, SessionMeta, read_json
from data.session_io import ImageEntry, save_session, set_image_included
from data.workspace import Workspace


def _fields(values, **kw):
    return [FieldResult(field_id=f"f{i}", G=v, method="planimetric", **kw)
            for i, v in enumerate(values)]


# ---------------------------------------------------------------- 1
def test_ci_matches_student_t():
    st = sample_statistics(_fields([7.0, 7.2, 7.4, 7.6, 7.8]), StatsConfig())
    assert st.n_fields == 5
    assert st.G_mean == pytest.approx(7.4, abs=1e-3)
    assert st.G_std == pytest.approx(0.3162, abs=1e-3)
    assert st.t_value == pytest.approx(2.776, abs=1e-3)
    assert st.G_ci95 == pytest.approx(0.3926, abs=1e-3)
    # E112 practice: %RA on N_A, G bounds from the N_A CI bounds
    assert st.basis == "N_A" and st.RA_pct is not None and st.RA_pct > 0
    assert st.G_ci_low < st.G_mean < st.G_ci_high
    assert "±" in st.summary_text()


def test_intercept_basis_bounds_are_ordered():
    fields = [FieldResult(field_id=f"f{i}", G=g, method="intercept")
              for i, g in enumerate([6.8, 7.0, 7.1, 7.3, 7.2])]
    st = sample_statistics(fields, {"required_fields": 5, "target_RA_pct": 10})
    assert st.basis == "l_bar" and st.method == "intercept"
    assert st.G_ci_low < st.G_mean < st.G_ci_high


# ---------------------------------------------------------------- 2
def test_single_field_ci_is_na_and_needs_more_fields():
    st = sample_statistics(_fields([7.4]), StatsConfig(required_fields=5))
    assert st.G_ci95 is None and st.RA_pct is None
    assert st.adequate is False
    assert st.status == "need ≥ 4 more fields" and st.status_level == "amber"


def test_uncalibrated_lot_is_grey():
    st = sample_statistics([FieldResult(field_id="a", G=None)] * 3)
    assert st.status == "uncalibrated" and st.status_level == "grey"
    assert st.n_fields == 0 and st.n_uncalibrated == 3


def test_adequate_lot_and_outlier_hint():
    vals = [7.40, 7.41, 7.39, 7.40, 7.42, 7.38, 7.40, 7.41, 7.39, 9.5]
    st = sample_statistics(_fields(vals))
    assert st.outlier_field_ids == ["f9"]
    st2 = sample_statistics(_fields(vals[:-1]))
    assert st2.adequate and st2.status == "adequate" and st2.status_level == "green"


# ---------------------------------------------------------------- 4
def test_fields_needed_iterates_t():
    n = fields_needed(0.1, 10.0)
    assert n >= 6
    # fixed point: n satisfies the E112 relation, n-1 does not
    from scipy import stats
    assert n >= (stats.t.ppf(0.975, n - 1) * 0.1 / 0.1) ** 2
    assert n - 1 < (stats.t.ppf(0.975, n - 2) * 0.1 / 0.1) ** 2
    # through sample_statistics: N_A values with s/x_bar = 0.1
    na = np.array([900.0, 1000.0, 1100.0, 1000.0, 900.0, 1100.0])
    na = na.mean() + (na - na.mean()) * (0.1 * na.mean() / na.std(ddof=1))
    Gs = [3.321928 * math.log10(x) - 2.954 for x in na]
    st = sample_statistics(_fields(Gs), StatsConfig(required_fields=2, target_RA_pct=10))
    assert st.n_needed == n >= 6


# ---------------------------------------------------------------- data helpers
def _result(G, ecd=2.0, n_grains=4):
    grains = [GrainResult(grain_id=i + 1, area_px=100.0, area_um2=math.pi * ecd ** 2 / 4,
                          perimeter_px=40.0, perimeter_um=5.0, equivalent_diameter_px=11.3,
                          equivalent_diameter_um=ecd, major_axis_um=ecd, minor_axis_um=ecd,
                          aspect_ratio=1.0, circularity=0.9, eccentricity=0.1,
                          centroid_x=5.0, centroid_y=5.0, bbox=(0, 0, 10, 10))
              for i in range(n_grains)]
    r = AnalysisResult()
    r.grains = grains
    r.grain_count = n_grains
    r.px_per_um = 5.0
    r.has_calibration = True
    r.mean_diameter_um = ecd
    r.astm_g = G
    r.astm = {"primary_method": "planimetric", "G_primary": G}
    return r


def _lot(tmp_path, sessions):
    ws = Workspace(tmp_path / "ws")
    proj = ws.create_project("P1")
    sample = ws.create_sample(proj, "S1")
    lot = ws.create_lot(proj, sample, "L1")
    img = np.full((32, 32, 3), 128, np.uint8)
    refs = []
    for label, gs in sessions:
        entries = [ImageEntry(image_bgr=img, result=_result(g), filename=f"{label}_{i}.png")
                   for i, g in enumerate(gs)]
        refs.append(save_session(lot, {"project": "P1", "sample_id": "S1", "lot_number": "L1"},
                                 entries, label=label))
    return ws, lot, refs


# ---------------------------------------------------------------- 3
def test_exclusion_changes_n_and_is_audited(tmp_path):
    _, lot, refs = _lot(tmp_path, [("A", [7.0, 7.2, 7.4, 7.6, 7.8])])
    assert sample_statistics(fields_for_lot(lot)).n_fields == 5
    with pytest.raises(ValueError):
        set_image_included(refs[0].path, "A_2.png", False, reason="  ")
    note = set_image_included(refs[0].path, "A_2.png", False, reason="scratch in field",
                              operator="jack")
    fields = fields_for_lot(lot)
    st = sample_statistics(fields)
    assert st.n_fields == 4 and st.n_excluded == 1
    ex = [f for f in fields if not f.included]
    assert len(ex) == 1 and ex[0].exclusion_reason == "scratch in field"
    m = read_json(refs[0].path / "manifest.json")
    assert m["audit_log"][-1]["action"] == "exclude_field"
    assert m["audit_log"][-1]["reason"] == "scratch in field"
    assert m["audit_log"][-1]["operator"] == "jack" and note["image"] == "A_2.png"
    set_image_included(refs[0].path, "A_2.png", True)
    assert sample_statistics(fields_for_lot(lot)).n_fields == 5
    assert len(read_json(refs[0].path / "manifest.json")["audit_log"]) == 2


def test_old_manifest_loads_as_included():
    e = ImageManifestEntry.from_dict({"filename": "x.png", "has_result": True})
    assert e.included is True and e.exclusion_reason is None
    m = SessionMeta.from_dict({"session_id": "s", "images": [{"filename": "x.png"}]})
    assert m.audit_log == [] and m.images[0].included
    s = AppSettings.from_dict({"theme": "dark"})
    assert s.required_fields == 5 and s.target_RA_pct == 10.0


# ---------------------------------------------------------------- 5
def test_lot_pools_sessions_session_scope_uses_one(tmp_path):
    ws, lot, refs = _lot(tmp_path, [("A", [7.0, 7.2, 7.4]), ("B", [7.6, 7.8])])
    fields = fields_for_lot(lot)
    assert len(fields) == 5
    st = sample_statistics(fields, AppSettings())
    assert st.n_fields == 5 and st.G_mean == pytest.approx(7.4)
    assert st.G_ci95 == pytest.approx(0.3926, abs=1e-3)
    sid_b = refs[1].session_id
    st_b = sample_statistics(fields, {"scope": "session", "session_id": sid_b})
    assert st_b.n_fields == 2 and st_b.G_mean == pytest.approx(7.7)
    assert st.ecd_mean_um == pytest.approx(2.0)
    assert len(pooled_ecds_um(fields)) == 5 * 4
    # catalog entry point: by folder and by lot number
    cat = Catalog(ws.root)
    cat.rebuild()
    assert len(cat.fields_for_lot(lot)) == 5
    assert len(cat.fields_for_lot("L1")) == 5


# ---------------------------------------------------------------- 6
def test_reports_lot_block_and_pptx_tile(tmp_path):
    import cv2
    from pptx import Presentation
    from reports.excel_renderer import render_excel
    from reports.model import ReportImageInput, ReportModel
    from reports.pptx_renderer import render_pptx

    items = []
    for i, g in enumerate([7.0, 7.2, 7.4, 7.6, 7.8]):
        p = str(tmp_path / f"f{i}.png")
        cv2.imwrite(p, np.full((32, 32, 3), 128, np.uint8))
        items.append(ReportImageInput(image_path=p, result=_result(g), lot_number="L1"))
    model = ReportModel.from_results(items, title="Lot", sample_statistics="auto",
                                     asset_dir=str(tmp_path / "assets"))
    assert len(model.sample_statistics) == 1
    assert model.sample_statistics[0]["G_ci95"] == pytest.approx(0.3926, abs=1e-3)
    # JSON round-trip keeps the section
    assert ReportModel.from_json(model.to_json()).sample_statistics == model.sample_statistics

    xlsx = render_excel(model, str(tmp_path / "r.xlsx"))
    ws = openpyxl.load_workbook(xlsx)["Overview"]
    cells = {(c.row, c.column): c for row in ws.iter_rows() for c in row if c.value is not None}
    hdr = {c.value: (r, col) for (r, col), c in cells.items() if isinstance(c.value, str)}
    assert "± 95% CI (G)" in hdr and "%RA" in hdr and "Lot statistics" in \
        next(v for v in hdr if v.startswith("Lot statistics"))
    r, col = hdr["± 95% CI (G)"]
    ci_cell = ws.cell(row=r + 1, column=col)
    assert ci_cell.value == pytest.approx(0.3926, abs=1e-3)
    assert ci_cell.number_format == "0.00"
    r, col = hdr["%RA"]
    ra_cell = ws.cell(row=r + 1, column=col)
    assert isinstance(ra_cell.value, float) and ra_cell.number_format == "0.00"
    # lot block sits above the per-image table
    assert hdr["%RA"][0] < hdr["Grains"][0]

    pptx = render_pptx(model, str(tmp_path / "r.pptx"))
    texts = [sh.text_frame.text for s in Presentation(pptx).slides for sh in s.shapes
             if sh.has_text_frame]
    assert any(t.startswith("G 7.40 ± 0.39") for t in texts)


def test_reports_unchanged_without_statistics(tmp_path):
    from reports.model import ReportModel
    m = ReportModel()
    assert m.sample_statistics == []
    assert ReportModel.from_dict({"title": "x"}).sample_statistics == []
