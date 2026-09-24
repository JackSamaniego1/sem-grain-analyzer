"""Tests for reports.pptx_renderer.render_pptx: slide count/titles/native charts."""
import os
import sys
import tempfile

import numpy as np
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import cv2

from tests.conftest import make_mosaic
from core.grain_detector import GrainDetector, DetectionParams
import pytest

from reports.model import ReportModel, ReportImageInput, Section
from reports.pptx_renderer import render_pptx
from reports.charts import PALETTES


def _build_model(tmp_path, n=3, px_per_um=8.0, seeds=None):
    det = GrainDetector()
    seeds = seeds or list(range(1, n + 1))
    items = []
    for i, seed in enumerate(seeds):
        gray, _ = make_mosaic(seed=seed, h=256, w=256, n_grains=30)
        bgr = np.repeat(gray[:, :, None], 3, axis=2)
        res = det.analyze(bgr, px_per_um=px_per_um, params=DetectionParams())
        img_path = str(tmp_path / f"synth_{i}.png")
        cv2.imwrite(img_path, bgr)
        items.append(ReportImageInput(
            image_path=img_path, result=res, image_bgr=bgr,
            sample_id=f"S{i}", lot_number="L1",
        ))
    return ReportModel.from_results(
        items, title="Test Deck", operator="Jack", organization="Acme",
        metadata={"detection_mode": "boundary", "detection_params": {"min_grain_size_px": 50}},
        asset_dir=str(tmp_path / "assets"),
    )


def _expected_slide_count(n_images, want_charts=True, want_methods=True):
    return 1 + 1 + (2 if want_charts else 0) + n_images + (1 if want_methods else 0) + 1


def _build_mixed_model(tmp_path):
    """One calibrated image + one uncalibrated image."""
    det = GrainDetector()
    items = []
    for i, (seed, ppu) in enumerate([(1, 8.0), (2, 0.0)]):
        gray, _ = make_mosaic(seed=seed, h=256, w=256, n_grains=30)
        bgr = np.repeat(gray[:, :, None], 3, axis=2)
        res = det.analyze(bgr, px_per_um=ppu, params=DetectionParams())
        img_path = str(tmp_path / f"synth_{i}.png")
        cv2.imwrite(img_path, bgr)
        items.append(ReportImageInput(image_path=img_path, result=res, image_bgr=bgr,
                                       sample_id=f"S{i}", lot_number="L1"))
    return ReportModel.from_results(items, title="Mixed", asset_dir=str(tmp_path / "assets"))


def _cell_value(text):
    """Extract the leading float from a 'NN.NN unit' table/callout string."""
    return float(text.split()[0])


def test_uncalibrated_exec_summary_row_has_no_zero_size_stats(tmp_path):
    model = _build_model(tmp_path, n=1, px_per_um=0.0)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    table = [sh for sh in prs.slides[1].shapes if sh.has_table][0].table
    row = table.rows[1]
    mean_diam_text = table.cell(1, 2).text
    mean_area_text = table.cell(1, 4).text
    assert _cell_value(mean_diam_text) > 0
    assert _cell_value(mean_area_text) > 0
    assert "px" in mean_diam_text and "px" in mean_area_text


def test_mixed_calibration_exec_summary_rows_have_correct_units_and_no_zeros(tmp_path):
    model = _build_mixed_model(tmp_path)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    table = [sh for sh in prs.slides[1].shapes if sh.has_table][0].table
    calibrated_row_text = [table.cell(1, c).text for c in range(len(table.columns))]
    uncal_row_text = [table.cell(2, c).text for c in range(len(table.columns))]
    assert _cell_value(calibrated_row_text[2]) > 0 and "µm" in calibrated_row_text[2]
    assert _cell_value(uncal_row_text[2]) > 0 and "px" in uncal_row_text[2]
    assert _cell_value(calibrated_row_text[4]) > 0
    assert _cell_value(uncal_row_text[4]) > 0  # was 0.00 before the fix


def test_uncalibrated_image_slide_metric_callouts_have_no_zeros(tmp_path):
    model = _build_model(tmp_path, n=1, px_per_um=0.0)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    image_slide = prs.slides[4]  # title, exec, area-dist, diam-dist, img1
    texts = [
        run.text for shape in image_slide.shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs for run in p.runs
    ]
    joined = " ".join(texts)
    assert "px" in joined
    # Mean diameter / mean area callouts should not both read as 0.00.
    zero_like = [t for t in texts if t.strip() in ("0.00 px", "0.00 px²")]
    assert zero_like == []


def test_slide_count_for_n_images(tmp_path):
    model = _build_model(tmp_path, n=3)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(3)


def test_is_16_9(tmp_path):
    model = _build_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    ratio = prs.slide_width / prs.slide_height
    assert abs(ratio - 16 / 9) < 0.01


def test_title_slide_has_report_title(tmp_path):
    model = _build_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    slide0_text = "\n".join(
        run.text for shape in prs.slides[0].shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs for run in p.runs
    )
    assert "Test Deck" in slide0_text


def test_executive_summary_table_has_one_row_per_image_plus_combined(tmp_path):
    model = _build_model(tmp_path, n=3)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    exec_slide = prs.slides[1]
    tables = [sh for sh in exec_slide.shapes if sh.has_table]
    assert len(tables) == 1
    table = tables[0].table
    assert len(table.rows) == 3 + 2  # header + 3 images + combined
    assert table.cell(len(table.rows) - 1, 0).text == "Combined"


def test_distribution_slides_have_native_charts_with_axis_titles(tmp_path):
    model = _build_model(tmp_path, n=2)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    dist_slides = list(prs.slides)[2:4]
    units_seen = []
    for slide in dist_slides:
        charts = [sh for sh in slide.shapes if sh.has_chart]
        assert len(charts) == 1
        chart = charts[0].chart
        cat_title = chart.category_axis.axis_title.text_frame.text
        val_title = chart.value_axis.axis_title.text_frame.text
        assert val_title == "Number of Grains"
        assert "(" in cat_title and ")" in cat_title  # units present
        units_seen.append(cat_title)
    assert any("Area" in u for u in units_seen)
    assert any("Diameter" in u for u in units_seen)


def test_image_slides_have_metric_callouts_and_caption(tmp_path):
    model = _build_model(tmp_path, n=2)
    model.images[0].caption = "Notably coarse grains"
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    image_slide = prs.slides[4]  # title, exec, area, diam, img1
    text = "\n".join(
        run.text for shape in image_slide.shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs for run in p.runs
    )
    assert "Grains" in text
    assert "Mean Diameter" in text
    assert "Notably coarse grains" in text


def test_methods_and_appendix_slides_present(tmp_path):
    model = _build_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    joined = "\n".join(all_text)
    assert "Methods" in joined
    assert "Excel" in joined  # appendix note


def test_disabling_combined_distribution_removes_those_slides(tmp_path):
    model = _build_model(tmp_path, n=2)
    model.get_section("combined_distribution").enabled = False
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(2, want_charts=False)


def test_excluding_image_reduces_slide_count(tmp_path):
    model = _build_model(tmp_path, n=3)
    model.images[0].include = False
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(2)


def test_no_temp_files_leaked(tmp_path):
    model = _build_model(tmp_path, n=2)
    tmpdir = tempfile.gettempdir()
    before = set(os.listdir(tmpdir))
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    after = set(os.listdir(tmpdir))
    leaked = [n for n in (after - before) if "grain_report_pptx" in n]
    assert leaked == []


def _all_text(slide):
    return "\n".join(
        run.text for shape in slide.shapes if shape.has_text_frame
        for p in shape.text_frame.paragraphs for run in p.runs)


def _add_custom_text(model, order, title, body):
    sec = Section(id=f"text_{title}", type="custom_text", title=title, enabled=True,
                  order=order, payload={"body": body})
    model.sections.append(sec)
    return sec


# ---------------------------------------------------------------------------
# REP-08: designer edits the renderer must honour
# ---------------------------------------------------------------------------

def test_reordering_top_level_sections_reorders_slides(tmp_path):
    """Move Methods before the distribution slides purely via ``Section.order``."""
    model = _build_model(tmp_path, n=1)
    model.get_section("parameters").order = 1.5   # between overview(1) and charts(2)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    titles = [_all_text(s).split("\n")[0] for s in prs.slides]
    methods_idx = next(i for i, t in enumerate(titles) if "Methods" in t)
    area_idx = next(i for i, t in enumerate(titles) if "Area Distribution" in t)
    assert methods_idx < area_idx
    # Appendix (raw data) is still strictly last.
    assert "Appendix" in titles[-1]


def test_renamed_custom_text_title_used_as_slide_heading(tmp_path):
    model = _build_model(tmp_path, n=1)
    sec = _add_custom_text(model, 1.1, "Sample Preparation", "Etched per ASTM E407.")
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    joined = "\n".join(_all_text(s) for s in prs.slides)
    assert "Sample Preparation" in joined
    assert "Etched per ASTM E407." in joined


def test_disabling_cover_removes_title_slide(tmp_path):
    model = _build_model(tmp_path, n=1)
    model.get_section("cover").enabled = False
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(1) - 1
    assert "Executive Summary" in _all_text(prs.slides[0])


def test_disabling_overview_table_removes_exec_summary_slide(tmp_path):
    model = _build_model(tmp_path, n=1)
    model.get_section("overview_table").enabled = False
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(1) - 1
    assert not any(sh.has_table for sh in prs.slides[1].shapes)


@pytest.mark.parametrize("theme_id", list(PALETTES.keys()))
def test_each_palette_recolours_title_slide(tmp_path, theme_id):
    from pptx.dml.color import RGBColor
    model = _build_model(tmp_path, n=1)
    model.theme = theme_id
    out = str(tmp_path / f"deck_{theme_id}.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    title_shape = next(sh for sh in prs.slides[0].shapes if sh.has_text_frame
                        and model.title in sh.text_frame.text)
    run = title_shape.text_frame.paragraphs[0].runs[0]
    expected = RGBColor.from_string(PALETTES[theme_id]["accent"].lstrip("#"))
    assert run.font.color.rgb == expected


def test_two_custom_text_sections_render_native_text_slides_in_order(tmp_path):
    model = _build_model(tmp_path, n=1)
    _add_custom_text(model, 1.1, "Sample Prep", "Etched per ASTM E407.")
    _add_custom_text(model, 1.2, "Acceptance Criteria", "Grain size must be G >= 5.")
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    titles = [_all_text(s).split("\n")[0] for s in prs.slides]
    assert titles.count("") == 0
    prep_idx = titles.index("Sample Prep")
    accept_idx = titles.index("Acceptance Criteria")
    exec_idx = titles.index("Executive Summary")
    area_idx = next(i for i, t in enumerate(titles) if "Area Distribution" in t)
    assert exec_idx < prep_idx < accept_idx < area_idx
    assert len(prs.slides) == _expected_slide_count(1) + 2


def test_disabled_custom_text_section_is_not_rendered(tmp_path):
    model = _build_model(tmp_path, n=1)
    sec = _add_custom_text(model, 1.1, "Draft", "not ready")
    sec.enabled = False
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    assert len(prs.slides) == _expected_slide_count(1)
    assert "Draft" not in "\n".join(_all_text(s) for s in prs.slides)


# ---------------------------------------------------------------------------
# HIER-01: user-defined folder hierarchy in the PowerPoint renderer
# ---------------------------------------------------------------------------

_HIERARCHY = [
    {"key": "project", "label": "Job #", "value": "24-117"},
    {"key": "sample", "label": "Part Number", "value": "7718-A"},
    {"key": "lot", "label": "Lot", "value": "L-44A"},
]


def _build_hierarchy_model(tmp_path, n=2, display_names=None):
    det = GrainDetector()
    items = []
    for i in range(n):
        gray, _ = make_mosaic(seed=i + 1, h=256, w=256, n_grains=30)
        bgr = np.repeat(gray[:, :, None], 3, axis=2)
        res = det.analyze(bgr, px_per_um=8.0, params=DetectionParams())
        img_path = str(tmp_path / f"src_{i}.png")
        cv2.imwrite(img_path, bgr)
        item = ReportImageInput(image_path=img_path, result=res, image_bgr=bgr)
        if display_names:
            item.display_name = display_names[i]
        items.append(item)
    return ReportModel.from_results(items, title="Job Report", hierarchy=_HIERARCHY,
                                     asset_dir=str(tmp_path / "assets"))


def test_title_slide_shows_hierarchy_lines(tmp_path):
    model = _build_hierarchy_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    text = _all_text(prs.slides[0])
    assert "Job #: 24-117" in text
    assert "Part Number: 7718-A" in text
    assert "Lot: L-44A" in text


def test_footer_shows_hierarchy_and_page_number(tmp_path):
    model = _build_hierarchy_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    text = _all_text(prs.slides[0])
    assert "Job # 24-117" in text
    assert "Part Number 7718-A" in text
    assert "Lot L-44A" in text
    assert "page 1" in text


def test_image_slide_title_uses_display_name(tmp_path):
    model = _build_hierarchy_model(tmp_path, n=1, display_names=["L-44A_01"])
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    image_slide = prs.slides[4]  # title, exec, area, diam, img1
    text = _all_text(image_slide)
    assert "L-44A_01" in text
    assert "src_0" not in text


def test_exec_summary_table_uses_hierarchy_column_labels(tmp_path):
    model = _build_hierarchy_model(tmp_path, n=1, display_names=["L-44A_01"])
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    table = [sh for sh in prs.slides[1].shapes if sh.has_table][0].table
    headers = [table.cell(0, c).text for c in range(len(table.columns))]
    assert "Job #" in headers and "Part Number" in headers and "Lot" in headers
    row1 = [table.cell(1, c).text for c in range(len(table.columns))]
    assert "L-44A_01" in row1
    assert "24-117" in row1 and "7718-A" in row1 and "L-44A" in row1


def test_legacy_model_without_hierarchy_keeps_old_footer_and_title(tmp_path):
    model = _build_model(tmp_path, n=1)
    out = str(tmp_path / "deck.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    text = _all_text(prs.slides[0])
    assert "Sample/Lot:" in text
    assert model.title in text
    assert "page 1" not in text  # legacy footer has no "page n" phrase


def test_sample_output_written_to_scratch():
    scratch = os.path.join(ROOT, "scratch", "reports")
    os.makedirs(scratch, exist_ok=True)
    det = GrainDetector()
    items = []
    for i, seed in enumerate([21, 22, 23]):
        gray, _ = make_mosaic(seed=seed, h=256, w=256, n_grains=30)
        bgr = np.repeat(gray[:, :, None], 3, axis=2)
        res = det.analyze(bgr, px_per_um=8.0, params=DetectionParams())
        img_path = os.path.join(scratch, f"sample_src2_{i}.png")
        cv2.imwrite(img_path, bgr)
        items.append(ReportImageInput(image_path=img_path, result=res, image_bgr=bgr,
                                       sample_id=f"S{i}", lot_number="L1"))
    model = ReportModel.from_results(
        items, title="Sample Grain Report", operator="Jack", organization="Acme",
        metadata={"detection_mode": "boundary"}, asset_dir=os.path.join(scratch, "assets2"),
    )
    model.theme = "slate_teal"
    _add_custom_text(model, 1.1, "Sample Preparation", "Mounted, polished, etched per ASTM E407.")
    _add_custom_text(model, 20_001, "Conclusions", "Grain size meets the acceptance criterion.")
    out = os.path.join(scratch, "sample.pptx")
    render_pptx(model, out)
    assert os.path.exists(out)
    prs = Presentation(out)
    joined = "\n".join(_all_text(s) for s in prs.slides)
    assert "Sample Preparation" in joined
    assert "Conclusions" in joined
