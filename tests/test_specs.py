"""INN-02: spec limits & PASS/FAIL/INCONCLUSIVE conformity verdict
(ILAC-G8:09/2019 guarded/simple acceptance).

Spec limits are strictly opt-in: with no ``Spec`` attached, nothing should
appear anywhere -- no verdict, no badge, no Excel/PPTX text, no extra
``lot_verdicts`` cache rows (see the last section of this file)."""
import os
import sys

import numpy as np
import openpyxl
import pytest
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

import cv2

from core.grain_detector import AnalysisResult, GrainResult
from core.metrics import FieldResult, SampleStatistics, sample_statistics
from data.catalog import Catalog, compute_lot_verdict, fields_for_lot
from data.models import ImageEntry
from data.session_io import save_session
from data.specs import (FAIL, INCONCLUSIVE, NO_SPEC, PASS, Rule, Spec, SpecError, Verdict,
                        evaluate, known_metrics, select_spec, specs_from_project_dict,
                        specs_to_project_dict)
from data.workspace import Workspace
from reports.model import ReportImageInput, ReportModel, normalize_verdict
from reports.excel_renderer import render_excel
from reports.pptx_renderer import render_pptx
from tests.conftest import make_mosaic


def _fields(values, **kw):
    return [FieldResult(field_id=f"f{i}", G=v, method="planimetric", **kw)
            for i, v in enumerate(values)]


# ---------------------------------------------------------------- 1
# "G 7.4 CI [7.2, 7.6], limits 6-8 -> PASS; limits 7.5-9 -> INCONCLUSIVE
# (guarded), FAIL (simple, mean 7.4 < 7.5); limits 8-9 -> FAIL."
def _g_stats():
    return SampleStatistics(G_mean=7.4, G_ci_low=7.2, G_ci_high=7.6, n_fields=5, adequate=True)


def test_guarded_pass_within_ci():
    spec = Spec(name="Alloy 718", revision="A", decision_rule="guarded",
                rules=[Rule(metric="G_mean", lower=6.0, upper=8.0)])
    v = evaluate(spec, _g_stats())
    assert v.overall == PASS
    assert v.rules[0].status == PASS


def test_guarded_inconclusive_when_ci_straddles_limit():
    spec = Spec(name="Alloy 718", revision="A", decision_rule="guarded",
                rules=[Rule(metric="G_mean", lower=7.5, upper=9.0)])
    v = evaluate(spec, _g_stats())
    assert v.overall == INCONCLUSIVE


def test_simple_fails_on_mean_below_lower():
    spec = Spec(name="Alloy 718", revision="A", decision_rule="simple",
                rules=[Rule(metric="G_mean", lower=7.5, upper=9.0)])
    v = evaluate(spec, _g_stats())
    assert v.overall == FAIL


def test_guarded_fails_when_ci_fully_outside():
    spec = Spec(name="Alloy 718", revision="A", decision_rule="guarded",
                rules=[Rule(metric="G_mean", lower=8.0, upper=9.0)])
    v = evaluate(spec, _g_stats())
    assert v.overall == FAIL


# ---------------------------------------------------------------- 2
def test_n_fields_below_minimum_is_inconclusive_with_reason():
    spec = Spec(name="Min fields", revision="1", rules=[Rule(metric="n_fields", lower=5)])
    stats = SampleStatistics(n_fields=3, G_mean=7.4, G_ci_low=7.2, G_ci_high=7.6, adequate=False)
    v = evaluate(spec, stats)
    assert v.overall == INCONCLUSIVE
    rule = v.rules[0]
    assert rule.status == INCONCLUSIVE
    assert rule.text == "fields 3 < 5"


# ---------------------------------------------------------------- 3
def test_sample_level_override_beats_project_spec():
    project_spec = Spec(id="proj", name="Project default", rules=[Rule(metric="G_mean", lower=6, upper=8)])
    sample_spec = Spec(id="override", name="Product-form override",
                       rules=[Rule(metric="G_mean", lower=5, upper=9)],
                       applies_to={"sample_ids": ["S-2024-0917"]})
    specs = [project_spec, sample_spec]
    assert select_spec(specs, "S-2024-0917").id == "override"
    assert select_spec(specs, "S-other").id == "proj"
    assert select_spec([project_spec], "anything").id == "proj"
    assert select_spec([sample_spec], "no-match") is None


# ---------------------------------------------------------------- 4
def test_spec_round_trips_through_project_dict():
    spec = Spec(id="sp1", name="Alloy 718 bar", revision="Rev C", decision_rule="guarded",
                rules=[Rule(metric="G_mean", lower=6.0, upper=8.0),
                       Rule(metric="RA_pct", upper=10.0),
                       Rule(metric="n_fields", lower=5)],
                applies_to={"sample_ids": ["S1"]})
    project_dict = specs_to_project_dict({"name": "P1"}, [spec])
    restored = specs_from_project_dict(project_dict)
    assert len(restored) == 1
    r = restored[0]
    assert (r.id, r.name, r.revision, r.decision_rule) == ("sp1", "Alloy 718 bar", "Rev C", "guarded")
    assert [rule.metric for rule in r.rules] == ["G_mean", "RA_pct", "n_fields"]
    assert r.applies_to == {"sample_ids": ["S1"]}


def test_unknown_metric_raises_clear_error():
    spec = Spec(name="Bad spec", rules=[Rule(metric="bogus_metric", upper=10)])
    with pytest.raises(SpecError, match="bogus_metric"):
        evaluate(spec, _g_stats())
    assert "G_mean" in known_metrics()


def test_no_spec_gives_no_spec_verdict():
    v = evaluate(None, _g_stats())
    assert v.overall == NO_SPEC
    assert normalize_verdict(v) is None
    assert normalize_verdict(v.to_dict()) is None


# ---------------------------------------------------------------- 5 (XLSX / PPTX badges)
def _build_report_model(tmp_path, verdict=None):
    from core.grain_detector import GrainDetector, DetectionParams
    det = GrainDetector()
    gray, _ = make_mosaic(seed=1, h=128, w=128, n_grains=15)
    bgr = np.repeat(gray[:, :, None], 3, axis=2)
    res = det.analyze(bgr, px_per_um=8.0, params=DetectionParams())
    img_path = str(tmp_path / "synth_0.png")
    cv2.imwrite(img_path, bgr)
    item = ReportImageInput(image_path=img_path, result=res, image_bgr=bgr,
                            sample_id="S1", lot_number="L1")
    return ReportModel.from_results(
        [item], title="Conformity Test", operator="Jack",
        metadata={"detection_mode": "boundary", "detection_params": {}},
        asset_dir=str(tmp_path / "assets"), verdict=verdict,
    )


def _pass_verdict():
    return Verdict(overall=PASS, spec_id="sp1", spec_name="Alloy 718 bar", spec_revision="Rev C",
                   decision_rule="guarded",
                   rules=[__import__("data.specs", fromlist=["RuleResult"]).RuleResult(
                       metric="G_mean", value=7.4, ci_low=7.2, ci_high=7.6, lower=6.0, upper=8.0,
                       status=PASS, text="G 7.40 [7.20–7.60] within 6.00–8.00")],
                   statement="Conformity decided by guarded acceptance: the 95 % confidence "
                             "interval must lie within the specification (ILAC-G8:09/2019).")


def test_excel_verdict_cell_text_and_fill(tmp_path):
    model = _build_report_model(tmp_path, verdict=_pass_verdict())
    out = str(tmp_path / "out.xlsx")
    render_excel(model, out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Overview"]
    hit = None
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("Conformity:"):
                hit = cell
                break
        if hit:
            break
    assert hit is not None and "PASS" in hit.value
    fill_rgb = str(hit.fill.start_color.rgb or "").upper()
    assert fill_rgb.endswith("2E7D32")


def test_pptx_title_slide_contains_pass(tmp_path):
    model = _build_report_model(tmp_path, verdict=_pass_verdict())
    out = str(tmp_path / "out.pptx")
    render_pptx(model, out)
    prs = Presentation(out)
    title_slide = prs.slides[0]
    texts = [s.text_frame.text for s in title_slide.shapes if s.has_text_frame]
    assert any("PASS" in t for t in texts)


# ---------------------------------------------------------------- opt-in: no spec -> unchanged exports
def test_excel_and_pptx_unchanged_with_no_spec(tmp_path):
    model = _build_report_model(tmp_path, verdict=None)
    assert model.verdict is None

    xlsx_path = str(tmp_path / "nospec.xlsx")
    render_excel(model, xlsx_path)
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Overview"]
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Conformity:" not in all_text
    for word in ("PASS", "FAIL", "INCONCLUSIVE"):
        assert word not in all_text

    pptx_path = str(tmp_path / "nospec.pptx")
    render_pptx(model, pptx_path)
    prs = Presentation(pptx_path)
    title_slide = prs.slides[0]
    texts = [s.text_frame.text for s in title_slide.shapes if s.has_text_frame]
    assert not any(any(w in t for w in ("PASS", "FAIL", "INCONCLUSIVE")) for t in texts)
    # No extra shape (the badge) was added to the title slide.
    baseline_model = _build_report_model(tmp_path, verdict=None)
    baseline_path = str(tmp_path / "baseline.pptx")
    render_pptx(baseline_model, baseline_path)
    baseline_prs = Presentation(baseline_path)
    assert len(title_slide.shapes) == len(baseline_prs.slides[0].shapes)


def test_report_model_round_trip_with_no_verdict():
    m = ReportModel(title="t")
    assert m.verdict is None
    d = m.to_dict()
    assert d["verdict"] is None
    m2 = ReportModel.from_dict(d)
    assert m2.verdict is None
    # A verdict dict that only carries overall == "no_spec" also collapses.
    m3 = ReportModel.from_dict({**d, "verdict": {"overall": "no_spec"}})
    assert m3.verdict is None


# ---------------------------------------------------------------- 6 (catalog cache)
def _grain(i):
    return GrainResult(grain_id=i, area_px=100.0, area_um2=4.0, perimeter_px=40.0, perimeter_um=1.6,
                       equivalent_diameter_px=10.0, equivalent_diameter_um=2.0, major_axis_um=2.2,
                       minor_axis_um=1.8, aspect_ratio=1.2, circularity=0.9, eccentricity=0.5,
                       centroid_x=5.0, centroid_y=5.0, bbox=(0, 0, 10, 10))


def _result(G, ecd=2.0, n_grains=20):
    grains = [_grain(i) for i in range(n_grains)]
    r = AnalysisResult()
    r.grains = grains
    r.grain_count = n_grains
    r.px_per_um = 5.0
    r.has_calibration = True
    r.mean_diameter_um = ecd
    r.astm_g = G
    r.astm = {"primary_method": "planimetric", "G_primary": G}
    return r


def _lot_with_fields(tmp_path, gs, project="P1", sample="S1", lot="L1"):
    ws = Workspace(tmp_path / "ws")
    proj = ws.create_project(project)
    samp = ws.create_sample(proj, sample)
    lot_path = ws.create_lot(proj, samp, lot)
    img = np.full((32, 32, 3), 128, np.uint8)
    entries = [ImageEntry(image_bgr=img, result=_result(g), filename=f"f_{i}.png")
              for i, g in enumerate(gs)]
    save_session(lot_path, {"project": project, "sample_id": sample, "lot_number": lot}, entries)
    return ws, proj, lot_path


def test_catalog_rebuild_matches_live_evaluation(tmp_path):
    ws, proj, lot_path = _lot_with_fields(tmp_path, [7.40, 7.41, 7.39, 7.40, 7.42])
    spec = Spec(name="Alloy 718 bar", revision="A", decision_rule="guarded",
                rules=[Rule(metric="G_mean", lower=6.0, upper=8.0),
                       Rule(metric="n_fields", lower=5)])
    ws.update_project_meta(proj, specs=[spec.to_dict()])

    live = compute_lot_verdict(lot_path)
    assert live["overall"] == PASS

    cat = Catalog(ws.root)
    cat.rebuild()
    cached = cat.get_lot_verdict(lot_path)
    assert cached == live


def test_no_spec_leaves_lot_verdicts_table_empty(tmp_path):
    ws, proj, lot_path = _lot_with_fields(tmp_path, [7.0, 7.2, 7.4, 7.6, 7.8])
    cat = Catalog(ws.root)
    cat.rebuild()
    with cat._connection() as conn:
        n = conn.execute("SELECT COUNT(*) FROM lot_verdicts").fetchone()[0]
    assert n == 0
    got = cat.get_lot_verdict(lot_path)
    assert got["overall"] == NO_SPEC
