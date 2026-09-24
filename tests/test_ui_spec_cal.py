"""INN-02 spec-limits UI, INN-29 calibration-check UI, INN-27 settings.

Both features are OPTIONAL (user requirement): without a spec / with
calibration verification off, nothing new is shown and reports are
unchanged.  Offscreen; everything in tmp_path."""
import json
from pathlib import Path

import numpy as np
import pytest

pytest.importorskip("pytestqt")

from data.models import AppSettings, read_json  # noqa: E402
from data.settings import save_settings  # noqa: E402
from data.specs import Rule, Spec  # noqa: E402
from tests.test_ui_lot_result import _Toasts, _lot, _result, _show_lot  # noqa: E402


@pytest.fixture
def env(tmp_path, monkeypatch, qapp):
    from ui.design.theme import apply_theme, set_reduced_motion
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    root = tmp_path / "ws"
    save_settings(AppSettings(workspace_root=str(root), operator="Tester", theme="dark"))
    set_reduced_motion(True)
    apply_theme(qapp, "dark")
    yield root
    set_reduced_motion(False)


@pytest.fixture
def page(env, qtbot):
    from ui.app_state import AppState
    from ui.pages.projects_page import ProjectsPage
    w = ProjectsPage(AppState(), _Toasts())
    qtbot.addWidget(w)
    w.resize(1400, 1000)
    w.show()
    yield w
    w.close()


def _save_spec(project_dir: Path, *specs):
    from ui.dialogs.spec_editor_dialog import save_specs
    save_specs(project_dir, list(specs))


TIGHT = [7.3, 7.35, 7.4, 7.45, 7.5]


def _g_spec(lo, hi, rule="guarded", **kw):
    return Spec(id="s1", name="Alloy 718", revision="B", decision_rule=rule,
                rules=[Rule("G_mean", lo, hi)], **kw)


# ---------------------------------------------------------------- lot card badge
def test_no_spec_no_badge(page, env, qtbot):
    lot, _ = _lot(env, [("A", TIGHT)])
    lp = _show_lot(page, qtbot, lot)
    assert not lp.card.verdict_badge.isVisibleTo(lp.card)
    assert not lp.card.spec_line.isVisibleTo(lp.card)
    assert lp.verdict() is None


@pytest.mark.parametrize("lo,hi,rule,expect,kind", [
    (6.0, 9.0, "guarded", "PASS", "success"),
    (7.38, 9.0, "guarded", "INCONCLUSIVE", "warning"),
    (8.5, 9.0, "simple", "FAIL", "danger"),
])
def test_spec_badge_on_lot_card(page, env, qtbot, lo, hi, rule, expect, kind):
    lot, _ = _lot(env, [("A", TIGHT)])
    _save_spec(lot.parent.parent, _g_spec(lo, hi, rule))
    lp = _show_lot(page, qtbot, lot)
    b = lp.card.verdict_badge
    assert b.isVisibleTo(lp.card) and b.text() == expect and b.kind() == kind
    assert "Alloy 718 (B)" in b.toolTip() and "G" in b.toolTip()
    assert "ILAC-G8" in b.toolTip()
    assert lp.card.spec_line.isVisibleTo(lp.card)
    # verdict stays the whole lot's when the scope is one session
    assert lp.verdict().overall == expect.lower()


def test_sample_override_beats_project_spec(page, env, qtbot):
    lot, _ = _lot(env, [("A", TIGHT)])
    s_proj = _g_spec(6.0, 9.0)
    s_samp = _g_spec(8.5, 9.0)
    s_samp.id, s_samp.applies_to = "s2", {"sample_ids": ["S1"]}
    _save_spec(lot.parent.parent, s_proj, s_samp)
    lp = _show_lot(page, qtbot, lot)
    assert lp.card.verdict_badge.text() == "FAIL"


def test_invalid_spec_shows_neutral_badge_not_crash(page, env, qtbot):
    lot, _ = _lot(env, [("A", [7.0, 7.2])])
    pj = lot.parent.parent / "project.json"
    d = read_json(pj)
    d["specs"] = [{"id": "x", "name": "bad", "rules": [{"metric": "nope", "lower": 1}]}]
    pj.write_text(json.dumps(d), encoding="utf-8")
    lp = _show_lot(page, qtbot, lot)
    assert lp.card.verdict_badge.text() == "SPEC ERROR"
    assert lp.card.verdict_badge.kind() == "neutral"


# ---------------------------------------------------------------- editor
def test_menu_has_spec_action_for_project_and_sample(page, env):
    from ui.app_state import NodeRef
    lot, _ = _lot(env, [("A", [7.0])])
    proj, samp = lot.parent.parent, lot.parent
    texts = [a[0] for a in page.menu_actions(NodeRef("project", proj)) if a]
    assert "Spec limits (optional)…" in texts
    texts = [a[0] for a in page.menu_actions(NodeRef("sample", samp)) if a]
    assert any(t.startswith("Spec limits for this") for t in texts)
    texts = [a[0] for a in page.menu_actions(NodeRef("lot", lot)) if a]
    assert not any("Spec" in t for t in texts)


def test_editor_empty_state_add_validate_save(page, env, qtbot):
    from ui.app_state import NodeRef
    lot, _ = _lot(env, [("A", TIGHT)])
    proj = lot.parent.parent
    dlg = page.open_spec_editor(NodeRef("project", proj))
    qtbot.addWidget(dlg)
    assert dlg.right.currentIndex() == 0            # empty state
    dlg.add_spec()
    assert dlg.right.currentIndex() == 1
    # no name + G rule without limits -> refused, nothing written
    assert dlg.save() is False and dlg.error.isVisibleTo(dlg)
    assert not read_json(proj / "project.json").get("specs")
    dlg.name.setText("Bar")
    dlg.name.textEdited.emit("Bar")
    dlg.rules.cellWidget(0, 1).setText("6")
    dlg.rules.cellWidget(0, 2).setText("8")
    dlg.rules.cellWidget(0, 2).textEdited.emit("8")
    dlg.rule.set_current_index(1, animate=False)     # simple
    assert dlg.save() is True
    d = read_json(proj / "project.json")
    assert d["name"] == "P1"                          # other keys kept
    sp = d["specs"][0]
    assert sp["name"] == "Bar" and sp["decision_rule"] == "simple"
    assert sp["rules"][0] == {"metric": "G_mean", "lower": 6.0, "upper": 8.0, "unit": ""}
    assert sp["rules"][1]["metric"] == "n_fields" and sp["applies_to"] == {}
    assert page.toasts.has("Spec limits saved")


def test_editor_from_sample_defaults_to_override(page, env, qtbot):
    from ui.app_state import NodeRef
    lot, _ = _lot(env, [("A", [7.0])])
    dlg = page.open_spec_editor(NodeRef("sample", lot.parent))
    qtbot.addWidget(dlg)
    dlg.add_spec()
    assert dlg.scope.current_index() == 1
    assert dlg.specs[0].applies_to == {"sample_ids": ["S1"]}
    dlg.reject()


def test_editor_lower_above_upper_is_a_problem(env, qtbot):
    from ui.dialogs.spec_editor_dialog import SpecEditorDialog
    d = env / "P"
    d.mkdir(parents=True)
    (d / "project.json").write_text("{}", encoding="utf-8")
    dlg = SpecEditorDialog(d)
    qtbot.addWidget(dlg)
    probs = dlg.problems_for(Spec(name="x", rules=[Rule("G_mean", 9, 6)]))
    assert any("lower limit is above" in p for p in probs)


# ---------------------------------------------------------------- settings
@pytest.fixture
def settings_page(env, qtbot):
    from ui.app_state import AppState
    from ui.pages.settings_page import SettingsPage
    st = AppState()
    w = SettingsPage(st, _Toasts())
    qtbot.addWidget(w)
    w.resize(1200, 1400)
    w.show()
    return w


def test_settings_lot_statistics_controls(settings_page):
    w = settings_page
    assert w.req_fields.value() == 5 and w.target_ra.value() == 10.0
    w.req_fields.setValue(8)
    w.target_ra.setValue(7.5)
    assert w.state.settings.required_fields == 8
    assert w.state.settings.target_RA_pct == 7.5


def test_calibration_verification_off_by_default_and_hidden(settings_page):
    w = settings_page
    assert not w.cal_toggle.isChecked()
    assert not w.cal_box.isVisibleTo(w)
    w.cal_toggle.setChecked(True)
    assert w.state.settings.calibration_verification_enabled is True
    assert w.cal_box.isVisibleTo(w) and w.cal_btn.isVisibleTo(w)
    assert w._chips == []                     # no check recorded -> no "Cal due" nag


def _grating(period=20.0, n=256):
    x = np.arange(n, dtype=np.float64)
    row = 128 + 80 * np.sin(2 * np.pi * x / period)
    g = np.tile(row, (n, 1)).astype(np.uint8)
    return np.dstack([g, g, g])


def test_cal_check_dialog_measure_save_and_chip(settings_page, qtbot, tmp_path):
    import cv2
    w = settings_page
    w.cal_toggle.setChecked(True)
    img = tmp_path / "grating.png"
    cv2.imwrite(str(img), _grating(20.0))
    dlg = w.open_cal_check()
    qtbot.addWidget(dlg)
    assert dlg.preview.pix is None and not dlg.btn_save.isEnabled()
    dlg.s_name.setText("Grating A")
    dlg.s_pitch.setValue(2.0)
    dlg.s_cert.setText("SN 1234")
    assert dlg.add_standard() is not None and dlg.standard.count() == 1
    dlg.instrument.setCurrentText("Zeiss Sigma")
    dlg.mag.setText("1000")
    dlg.load_image(str(img))
    qtbot.waitUntil(lambda: dlg.ref is not None, timeout=10000)
    dlg.ppu.setValue(10.0)                    # 20 px period / 10 px/µm = 2.0 µm
    dlg.measure()
    qtbot.waitUntil(lambda: dlg.check is not None, timeout=20000)
    assert dlg.check.passed and abs(dlg.check.error_pct) < 0.5
    assert dlg.verdict.text() == "PASS" and dlg.preview.lines
    assert dlg.btn_save.isEnabled()
    dlg.save()
    qtbot.waitUntil(lambda: (w.state.root / "calibration" / "checks.jsonl").exists(),
                    timeout=10000)
    lines = (w.state.root / "calibration" / "checks.jsonl").read_text("utf-8").splitlines()
    assert len(lines) == 1 and json.loads(lines[0])["certificate_no"] == "SN 1234"
    qtbot.waitUntil(lambda: bool(w._chips), timeout=5000)
    assert w._chips[0].kind() == "success" and "Zeiss Sigma" in w._chips[0].text()
    assert any(i["name"] == "Zeiss Sigma" for i in w.state.settings.instruments)


def test_cal_check_manual_method(settings_page, qtbot, tmp_path):
    import cv2
    w = settings_page
    w.cal_toggle.setChecked(True)
    img = tmp_path / "g.png"
    cv2.imwrite(str(img), _grating(20.0))
    dlg = w.open_cal_check()
    qtbot.addWidget(dlg)
    dlg.s_name.setText("G")
    dlg.s_pitch.setValue(2.0)
    dlg.add_standard()
    dlg.instrument.setCurrentText("SEM-1")
    dlg.load_image(str(img))
    qtbot.waitUntil(lambda: dlg.ref is not None, timeout=10000)
    dlg.ppu.setValue(10.0)
    dlg.method.set_current_index(1, animate=False)
    assert not dlg.btn_measure.isVisibleTo(dlg) and dlg.n_periods.isVisibleTo(dlg)
    dlg.n_periods.setValue(10)
    dlg.preview.p0, dlg.preview.p1 = (10.0, 50.0), (215.0, 50.0)     # 205 px / 10
    dlg.preview.line_drawn.emit(dlg.preview.p0, dlg.preview.p1)       # = 20.5 px -> +2.5 %
    assert dlg.check is not None and dlg.check.method == "manual"
    assert dlg.check.error_pct == pytest.approx(2.5) and not dlg.check.passed
    assert dlg.verdict.text() == "FAIL"


def test_cal_chip_hidden_without_checks(env, qtbot):
    from data.cal_records import CalibrationStore
    from ui.dialogs.cal_check_dialog import CalStatusChip
    c = CalStatusChip()
    qtbot.addWidget(c)
    c.refresh(CalibrationStore(env, enabled=True), "SEM-1")
    assert c.isHidden() and c.state == "off"


# ---------------------------------------------------------------- report export
def _report_state(env, refs, **settings):
    class S:
        pass
    st = S()
    st.root = env
    st.settings = AppSettings(**settings)
    st.session = S()
    st.session.path = refs[0].path
    st.session.meta = S()
    st.session.meta.instrument = "SEM-1"
    st.session.meta.magnification = "1000"
    return st


def _inputs(refs, gs):
    from reports.model import ReportImageInput
    return [ReportImageInput(image_path=str(refs[0].path / "images" / f"A_{i}.png"),
                             result=_result(g), lot_number="L-1") for i, g in enumerate(gs)]


def test_report_unchanged_when_features_unused(env):
    from ui.pages import report_builder as rb
    lot, refs = _lot(env, [("A", [7.0, 7.4, 7.8])])
    st = _report_state(env, refs)
    inputs = _inputs(refs, [7.0, 7.4, 7.8])
    extras = rb.report_extras_arg(st, inputs)
    assert rb.compute_report_extras(extras) == (None, None)
    m = rb.build_model(inputs, title="t", operator="t", extras=extras)
    assert m.verdict is None and "calibration" not in m.metadata


def test_report_gets_verdict_and_calibration_when_used(env, tmp_path):
    from ui.pages import report_builder as rb
    lot, refs = _lot(env, [("A", TIGHT)])
    _save_spec(lot.parent.parent, _g_spec(6.0, 9.0))
    st = _report_state(env, refs, calibration_verification_enabled=True)
    inputs = _inputs(refs, TIGHT)
    m = rb.build_model(inputs, title="t", operator="t", extras=rb.report_extras_arg(st, inputs))
    assert m.verdict["overall"] == "pass" and m.verdict["spec_name"] == "Alloy 718"
    cal = m.metadata["calibration"]
    assert cal["status"] == "not verified" and cal["px_per_um"] == 5.0
    assert cal["text"].startswith("Scale not verified")
    # export-time recompute: spec removed -> verdict gone from the rendered files
    _save_spec(lot.parent.parent)
    out = tmp_path / "r.xlsx"
    rb.render_outputs(m.to_dict(), [("xlsx", str(out))],
                      extras=rb.report_extras_arg(_report_state(env, refs), inputs))
    assert out.exists()
    from openpyxl import load_workbook
    wb = load_workbook(out)
    txt = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows()
                   for c in row if c.value is not None)
    assert "PASS" not in txt
