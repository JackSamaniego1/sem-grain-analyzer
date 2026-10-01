"""PPTX contents page as slide 2 (UPDATE 4 item 19)."""
import os
import re
import sys

import numpy as np
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from reports.model import ImageSummary, ReportModel, Section
from reports.pptx_renderer import MAX_DATA_ROWS, render_pptx

HIER = [{"key": "sample", "label": "Part Number", "value": ""},
        {"key": "lot", "label": "Lot", "value": ""}]
ROW = re.compile(r"^(Page|Pages) (\d+)(?:–(\d+))?\s+·\s+(.+)$")


def _img(iid, part, lot, order):
    rng = np.random.default_rng(order)
    ds = rng.normal(3, 0.7, 40).clip(0.3)
    grains = [{"diameter_um": d, "area_um2": np.pi * (d / 2) ** 2,
               "diameter_px": d * 8, "area_px": np.pi * (d * 4) ** 2} for d in ds]
    return ImageSummary(id=iid, image_path="", order=order, px_per_um=8.0, has_calibration=True,
                        grain_count=len(grains), grains=grains, levels={"sample": part, "lot": lot})


def _model(parts=2, lots=4, per=2, n_text=0, contents=True):
    imgs, k = [], 0
    for p in range(parts):
        for l in range(lots):
            for _ in range(per):
                imgs.append(_img(f"i{k}", f"P{p + 1}", f"P{p + 1}-L{l + 1}", k))
                k += 1
    secs = [Section(id="cover", type="cover", title="Cover", order=0),
            Section(id="ov", type="overview_table", title="Overview", order=1),
            Section(id="ls", type="lot_summary", title="Lot Summary", order=500),
            Section(id="params", type="parameters", title="Methods", order=500)]
    for t in range(n_text):
        secs.append(Section(id=f"t{t}", type="custom_text", title=f"Note {t + 1}", order=2 + t,
                            payload={"body": "x"}))
    if not contents:
        secs.append(Section(id="c", type="contents", title="Contents", enabled=False, order=99))
    return ReportModel(title="Deck", sections=secs, images=imgs, hierarchy=HIER)


def _head(slide):
    return next(r.text for sh in slide.shapes if sh.has_text_frame
                for p in sh.text_frame.paragraphs for r in p.runs)


def _render(tmp_path, model):
    out = str(tmp_path / "d.pptx")
    render_pptx(model, out)
    return Presentation(out)


def _entries(prs):
    """[(first_page, last_page, label, target_slide_index0)] from every contents slide."""
    ids = [s.slide_id for s in prs.slides]
    res = []
    for s in prs.slides:
        if not _head(s).startswith("Contents"):
            continue
        for sh in s.shapes:
            if not sh.has_text_frame or sh.click_action.target_slide is None:
                continue
            m = ROW.match(sh.text_frame.text)
            assert m, sh.text_frame.text
            a = int(m.group(2))
            b = int(m.group(3) or a)
            res.append((a, b, m.group(4), ids.index(sh.click_action.target_slide.slide_id)))
    return res


def _footer_page(slide):
    return int([sh.text_frame.text for sh in slide.shapes if sh.has_text_frame][-1])


def test_contents_is_slide_2_and_summary_slide_3(tmp_path):
    prs = _render(tmp_path, _model())
    assert _head(prs.slides[0]) == "Deck"
    assert _head(prs.slides[1]) == "Contents"
    assert _head(prs.slides[2]) == "Grain Size Summary"


def test_every_link_targets_the_printed_page_and_matching_slide(tmp_path):
    prs = _render(tmp_path, _model())
    entries = _entries(prs)
    assert entries
    for first, last, label, target in entries:
        assert target == first - 1                    # link goes to the printed first page
        assert 1 <= first <= last <= len(prs.slides)
    # runs tile the deck (minus the contents slide itself) with no gaps
    covered = sorted(p for a, b, _l, _t in entries for p in range(a, b + 1))
    assert covered == [p for p in range(1, len(prs.slides) + 1) if p != 2]
    by_label = {label: (a, b, t) for a, b, label, t in entries}
    assert _head(prs.slides[by_label["Grain size summary"][2]]) == "Grain Size Summary"
    assert _head(prs.slides[by_label["Cover"][2]]) == "Deck"
    assert _head(prs.slides[by_label["Lot summary by part"][2]]).startswith("Lot Summary")
    assert _head(prs.slides[by_label["Grain distributions by part"][2]]).startswith("Grain Distributions")
    # removed sections never show up in the contents
    for gone in ("Appendix", "Methods & parameters", "Lot comparison by part",
                 "Lot-to-lot distribution comparison"):
        assert gone not in by_label
    # lot-chart slides come right after the summary
    assert _head(prs.slides[by_label["Lot summary, all lots"][2]]) == "Mean Grain Diameter by Lot"
    assert by_label["Lot summary, all lots"][0] == by_label["Grain size summary"][1] + 1
    assert by_label["Lot summary by part"][0] == by_label["Lot summary, all lots"][1] + 1


def test_similar_slides_collapse_to_ranges_pointing_at_first_slide(tmp_path):
    prs = _render(tmp_path, _model(parts=2, lots=4, per=2))   # 16 images, 8 lots
    by_label = {label: (a, b, t) for a, b, label, t in _entries(prs)}
    a, b, t = by_label["Image results"]
    assert b - a + 1 == 16 and t == a - 1
    assert _head(prs.slides[t]).startswith("Image")
    a, b, t = by_label["Grain distributions by part"]
    assert b - a + 1 == 2                                # one slide per part
    a, b, t = by_label["Lot summary by part"]
    assert b - a + 1 == 2 and _head(prs.slides[t]).startswith("Lot Summary")
    a, b, t = by_label["Image data tables"]
    assert b >= a and _head(prs.slides[t]).startswith("P1")
    # one line per run: the label appears exactly once
    labels = [e[2] for e in _entries(prs)]
    assert len(labels) == len(set(labels))


def test_long_deck_paginates_contents_and_numbers_stay_right(tmp_path):
    n_text = MAX_DATA_ROWS + 2                         # forces > 14 runs
    prs = _render(tmp_path, _model(parts=1, lots=2, per=1, n_text=n_text))
    heads = [_head(s) for s in prs.slides]
    assert heads[1] == "Contents" and heads[2] == "Contents"
    assert heads[3] == "Grain Size Summary"            # everything shifted by two
    entries = _entries(prs)
    assert len(entries) > MAX_DATA_ROWS
    for first, _last, label, target in entries:
        assert target == first - 1
    notes = [e for e in entries if e[2].startswith("Note ")]
    assert len(notes) == n_text
    for _a, _b, label, target in notes:
        assert heads[target] == label
    # footers carry the final page numbers
    assert [_footer_page(s) for s in prs.slides] == list(range(1, len(prs.slides) + 1))


def test_single_slide_entries_say_page_not_pages(tmp_path):
    prs = _render(tmp_path, _model(parts=1, lots=1, per=1))
    rows = [sh.text_frame.text for s in prs.slides if _head(s) == "Contents"
            for sh in s.shapes if sh.has_text_frame and ROW.match(sh.text_frame.text)]
    assert "Page 1  ·  Cover" in rows
    assert any(r.startswith("Pages ") for r in rows) or any(r.startswith("Page ") for r in rows)


def test_contents_disabled_gives_old_order(tmp_path):
    prs = _render(tmp_path, _model(contents=False))
    heads = [_head(s) for s in prs.slides]
    assert "Contents" not in heads
    assert heads[1] == "Grain Size Summary"
    assert [_footer_page(s) for s in prs.slides] == list(range(1, len(prs.slides) + 1))
    with_c = _render(tmp_path, _model())
    assert len(with_c.slides) == len(prs.slides) + 1


def test_cover_disabled_puts_contents_first(tmp_path):
    m = _model()
    m.get_section("cover").enabled = False
    prs = _render(tmp_path, m)
    assert _head(prs.slides[0]) == "Contents"
    assert _head(prs.slides[1]) == "Grain Size Summary"
    for first, _l, _lab, target in _entries(prs):
        assert target == first - 1
