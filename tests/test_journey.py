"""
REL-01 end-to-end user journeys (headless, pytest-qt/offscreen).

Each journey drives the real app shell/page objects (``ui.app_shell.AppShell``,
``ui.pages.analyze_page``/``review_page``/``reports_page``) the way a lab
operator would: open a project -> import images -> calibrate -> analyze with
a classic (non-SAM) detector -> edit grains -> save -> close and reopen the
project from disk -> build the report model -> export Excel (+PowerPoint) ->
reopen the exported files and check they carry the *edited* numbers.

Steps that would otherwise require a modal dialog (new-project wizard, a
scale-bar capture rectangle, a file-open dialog) go through the same
core/data APIs those dialogs call internally (``data.workspace.Workspace``,
``core.scale_bar.compute_px_per_um``, ``ui.app_state.AppState.probe_metadata``)
-- exactly the pattern the rest of this suite's shell tests already use
(see ``tests/ui_shell_helpers.py`` and ``tests/test_ui_shell_analysis.py``).

Detection uses the classic ``threshold``/``boundary`` pipelines only, so
these tests never need the (unbundled, 375 MB) SAM checkpoint -- see
``handoff/00_START_HERE.md`` "Risks". Nothing here touches the network
(``core/offline_guard.py`` is active for the whole suite via
``tests/test_offline.py``); all files are written under ``tmp_path``.
"""
from pathlib import Path

import openpyxl
import pytest
from pptx import Presentation

pytest.importorskip("pytestqt")

from data.models import AppSettings  # noqa: E402
from data.settings import save_settings  # noqa: E402
from tests.conftest import make_mosaic  # noqa: E402
from tests.ui_shell_helpers import make_session  # noqa: E402

TIMEOUT = 90000


# ======================================================================
# Shared fixtures / helpers (same idioms as test_ui_shell_analysis.py and
# test_ui_reports_designer.py)
# ======================================================================

@pytest.fixture
def env(tmp_path, monkeypatch, qapp):
    from ui.design.theme import apply_theme, set_reduced_motion
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    root = tmp_path / "ws"
    save_settings(AppSettings(workspace_root=str(root), operator="Tester", theme="dark"))
    set_reduced_motion(True)
    apply_theme(qapp, "dark")
    yield root
    set_reduced_motion(False)


def _open_shell(qtbot, path: Path, probe: bool = False):
    from ui.app_shell import AppShell
    from ui.app_state import AppState
    shell = AppShell(AppState(), probe_device=False)
    qtbot.addWidget(shell)
    shell.resize(1400, 900)
    shell.show()
    shell.open_session(path, probe=probe)
    qtbot.waitUntil(lambda: shell.state.session is not None, timeout=15000)
    return shell


def _analyse_all(shell, qtbot, mode="threshold"):
    st = shell.state
    shell.analyze.params.set_mode(mode)
    with qtbot.waitSignal(shell.analyze.queue.queue_finished, timeout=TIMEOUT):
        shell.analyze_all()
    qtbot.waitUntil(lambda: all(im.status == "done" and im.result is not None
                                for im in st.images()) and not st.is_filtering(),
                    timeout=TIMEOUT)


def _settle(shell, qtbot):
    st = shell.state
    qtbot.waitUntil(lambda: not st.is_filtering(), timeout=TIMEOUT)
    st.flush()
    qtbot.waitUntil(lambda: st.save_state == "saved" and not st.is_dirty(), timeout=TIMEOUT)


def _build_report(shell, qtbot):
    rp = shell.reports
    shell.go("reports")
    qtbot.waitUntil(lambda: not rp.is_busy(), timeout=TIMEOUT)
    rp.build_from_session()
    qtbot.waitUntil(lambda: rp.model is not None and not rp.is_busy(), timeout=TIMEOUT)
    return rp


def _export(rp, qtbot, kinds):
    with qtbot.waitSignal(rp.exported, timeout=TIMEOUT) as blocker:
        rp.export(list(kinds))
    return blocker.args[0]


def _overview_grain_counts(xlsx_path, n_images):
    """(per_image_counts, combined_total) read back from the exported
    Overview sheet's "Grains" column -- independent of whether the sheet
    used hierarchy columns or the Sample/Lot fallback."""
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb["Overview"]
    rows = list(ws.iter_rows(values_only=True))
    hdr_i = next(i for i, r in enumerate(rows) if r and "Grains" in r)
    col = rows[hdr_i].index("Grains")
    per_image = [rows[hdr_i + 1 + k][col] for k in range(n_images)]
    total = rows[hdr_i + 1 + n_images][col]
    return per_image, total


# ======================================================================
# Journey 1 (threshold mode): the full lifecycle, scale-bar calibration,
# trash + restore (undo/redo) grain edit, save/reload, report + both exports.
# ======================================================================

def test_journey_threshold_scale_bar_full_lifecycle(env, qtbot):
    from core.scale_bar import compute_px_per_um

    path = make_session(env, 2, label="Journey", project="AlloyCo",
                        sample="Bracket-A", lot="Lot-42")
    shell = _open_shell(qtbot, path)
    st = shell.state

    # ---- calibrate: scale-bar capture (a 100 px bar measured as 25 um) ----
    px_per_um = compute_px_per_um(100.0, 25.0)
    assert px_per_um == pytest.approx(4.0)
    st.set_calibration(px_per_um)
    assert st.session.px_per_um == pytest.approx(4.0)
    assert all(st.px_for(im) == pytest.approx(4.0) for im in st.images())

    # ---- analyze: classic threshold detector (no SAM checkpoint needed) ----
    _analyse_all(shell, qtbot, "threshold")
    base_counts = [im.result.grain_count for im in st.images()]
    assert all(c > 5 for c in base_counts)
    assert all(im.result.has_calibration for im in st.images())

    # ---- edit grains: trash a grain, restore it (undo), trash it again ----
    shell.go("review")
    im0 = st.images()[0]
    victim = im0.result.grains[0].grain_id
    shell.review.canvas.select([victim])
    shell.review.delete_selected()
    qtbot.waitUntil(lambda: not st.is_filtering()
                    and im0.result.grain_count == base_counts[0] - 1, timeout=TIMEOUT)
    assert victim in im0.manual and "manual" in im0.excluded[victim]
    assert victim in {g.grain_id for g in im0.raw.grains}          # non-destructive
    st.undo_stack.undo()                                            # restore from trash
    qtbot.waitUntil(lambda: not st.is_filtering()
                    and im0.result.grain_count == base_counts[0], timeout=TIMEOUT)
    assert victim not in im0.manual
    st.undo_stack.redo()                                            # trash it again, for the report
    qtbot.waitUntil(lambda: not st.is_filtering()
                    and im0.result.grain_count == base_counts[0] - 1, timeout=TIMEOUT)
    edited_counts = [base_counts[0] - 1, base_counts[1]]

    # ---- save ----
    _settle(shell, qtbot)
    shell.close()

    # ---- reload the project from disk ----
    shell2 = _open_shell(qtbot, path)
    st2 = shell2.state
    assert [im.result.grain_count for im in st2.images()] == edited_counts
    assert victim in st2.images()[0].manual
    assert st2.session.px_per_um == pytest.approx(4.0)

    # ---- build report model ----
    rp = _build_report(shell2, qtbot)
    m = rp.model
    ordered = sorted(m.images, key=lambda i: i.order)
    assert [i.grain_count for i in ordered] == edited_counts
    assert all(i.has_calibration for i in m.images)

    # ---- export Excel + PowerPoint, then reopen and check the numbers ----
    xlsx, pptx = _export(rp, qtbot, ["xlsx", "pptx"])
    assert Path(xlsx).is_file() and Path(pptx).is_file()

    per_image, total = _overview_grain_counts(xlsx, len(edited_counts))
    assert sorted(per_image) == sorted(edited_counts)
    assert total == sum(edited_counts)

    prs = Presentation(pptx)
    exec_table = next(sh for sh in prs.slides[1].shapes if sh.has_table).table
    headers = [exec_table.cell(0, c).text for c in range(len(exec_table.columns))]
    gcol = headers.index("Grains")
    pptx_counts = [int(exec_table.cell(r, gcol).text) for r in range(1, 1 + len(edited_counts))]
    assert sorted(pptx_counts) == sorted(edited_counts)

    shell2.close()


# ======================================================================
# Journey 2 (boundary mode + metadata calibration + merge edit)
# ======================================================================

# FEI/Thermo Fisher INI text block (real-world layout; see test_sem_metadata.py)
# carrying a vendor pixel-size field ("PixelWidth", metres/px) -> "high"
# confidence calibration that AppState.probe_metadata() auto-applies to a
# not-yet-analysed image.
_FEI_TEXT = (
    "[User]\r\nDate=03/12/2019\r\nTime=10:22:13 AM\r\nUser=supervisor\r\n\r\n"
    "[System]\r\nType=DualBeam\r\nDnumber=D9920\r\nSource=FEG\r\n"
    "Column=Elstar\r\nSystemType=Helios G4 UX\r\nDisplayWidth=0.4\r\n"
    "DisplayHeight=0.3\r\n\r\n[Beam]\r\nHV=5000\r\n\r\n[EBeam]\r\n"
    "Source=FEG\r\nHV=5000\r\nWD=0.00412\r\n\r\n[Scan]\r\n"
    "InternalScan=true\r\nDwelltime=3e-006\r\nPixelWidth=2.0345e-009\r\n"
    "PixelHeight=2.0345e-009\r\nHorFieldsize=2.0833e-006\r\n\r\n"
    "[Detectors]\r\nNumber=1\r\nName=TLD\r\nMode=SE\r\n"
)


def _make_fei_tiff_session(root: Path, tmp_path: Path) -> Path:
    tifffile = pytest.importorskip("tifffile")
    from data.catalog import Catalog
    from data.models import ImageEntry
    from data.session_io import save_session
    from data.workspace import Workspace

    src_dir = tmp_path / "src_tiff"
    src_dir.mkdir(parents=True, exist_ok=True)
    gray, _ = make_mosaic(seed=11, h=256, w=256, n_grains=40)
    tif_path = src_dir / "fei_0.tif"
    text = _FEI_TEXT.encode("latin-1")
    tifffile.imwrite(str(tif_path), gray, extratags=[(34682, 1, len(text), text, True)])

    ws = Workspace(root)
    pp = ws.create_project("MetaCalProj")
    sp = ws.create_sample(pp, "S-9", material="Steel")
    lp = ws.create_lot(pp, sp, "L-9")
    ref = save_session(lp, {"operator": "Tester"}, [ImageEntry(source_path=str(tif_path))],
                       label="MetaCal", catalog=Catalog(root))
    return Path(ref.path)


def test_journey_boundary_mode_metadata_calibration_and_merge(env, qtbot, tmp_path):
    from core.grain_edit import GrainEditError

    path = _make_fei_tiff_session(env, tmp_path)

    # ---- open with metadata probing on (the real "open from Projects" path) ----
    shell = _open_shell(qtbot, path, probe=True)
    st = shell.state
    im = st.images()[0]
    expected_ppu = 1.0 / 2.0345e-3       # 2.0345 nm/px vendor PixelWidth -> px/um
    qtbot.waitUntil(lambda: st.px_for(im) > 0, timeout=15000)
    assert st.px_for(im) == pytest.approx(expected_ppu, rel=1e-3)   # auto-applied, high confidence

    # ---- analyze: boundary detector (dense mosaic, thin dark grooves) ----
    _analyse_all(shell, qtbot, "boundary")
    assert im.result.grain_count > 5
    assert im.result.has_calibration

    # ---- edit grains: merge two neighbouring grains (falls back to a trash
    # if none of the closest pairs happen to touch within the merge gap) ----
    base = im.result.grain_count
    ids = [g.grain_id for g in im.result.grains]
    cent = {g.grain_id: (g.centroid_x, g.centroid_y) for g in im.result.grains}
    pairs = sorted(
        ((cent[a][0] - cent[b][0]) ** 2 + (cent[a][1] - cent[b][1]) ** 2, a, b)
        for i, a in enumerate(ids) for b in ids[i + 1:])
    merged = False
    for _, a, b in pairs[:20]:
        try:
            st.merge_grains(im.uid, [a, b])
            merged = True
            break
        except GrainEditError:
            continue
    if not merged:
        st.delete_grains(im.uid, [ids[0]])
    qtbot.waitUntil(lambda: not st.is_filtering()
                    and im.result.grain_count == base - 1, timeout=TIMEOUT)
    edited = base - 1

    # ---- save, close, reload from disk ----
    _settle(shell, qtbot)
    shell.close()
    shell2 = _open_shell(qtbot, path)
    st2 = shell2.state
    im2 = st2.images()[0]
    assert im2.result.grain_count == edited
    assert st2.px_for(im2) == pytest.approx(expected_ppu, rel=1e-3)

    # ---- build report + export Excel, check the edited count survived ----
    rp = _build_report(shell2, qtbot)
    assert rp.model.images[0].grain_count == edited

    xlsx, = _export(rp, qtbot, ["xlsx"])
    per_image, total = _overview_grain_counts(xlsx, 1)
    assert per_image[0] == edited and total == edited
    shell2.close()


# ======================================================================
# Journey 3 (negative): an uncalibrated image -> the report flags the
# missing calibration instead of silently showing bogus micron numbers.
# ======================================================================

def test_journey_uncalibrated_image_report_flags_missing_calibration(env, qtbot):
    path = make_session(env, 1, label="NoCal", project="RawProj",
                        sample="Coupon-1", lot="Lot-1")
    shell = _open_shell(qtbot, path)          # probe=False: calibration is never touched
    st = shell.state
    assert st.session.px_per_um == 0.0
    im = st.images()[0]
    assert im.px_override == 0.0

    _analyse_all(shell, qtbot, "threshold")
    assert im.result.grain_count > 5
    assert im.result.has_calibration is False

    _settle(shell, qtbot)
    rp = _build_report(shell, qtbot)
    m = rp.model
    assert all(not i.has_calibration for i in m.images)

    xlsx, = _export(rp, qtbot, ["xlsx"])
    wb = openpyxl.load_workbook(xlsx)

    methods_name = next(n for n in wb.sheetnames if "Method" in n)
    values = [c.value for row in wb[methods_name].iter_rows() for c in row if c.value is not None]
    assert values[values.index("Calibrated") + 1] == "No"

    per_image, total = _overview_grain_counts(xlsx, 1)
    assert per_image[0] == im.result.grain_count == total

    ov_text = " ".join(str(c.value) for row in wb["Overview"].iter_rows()
                       for c in row if c.value is not None)
    assert "px" in ov_text and "µm" not in ov_text        # units stayed in pixels, not microns
    shell.close()
