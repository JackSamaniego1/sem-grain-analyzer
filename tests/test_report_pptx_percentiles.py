"""D10/D50/D90 percentile slide (UPDATE 4 item 18), synthetic grains only."""
import os
import sys

import numpy as np
from pptx import Presentation

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from reports.model import ReportModel, ImageSummary, Section, pooled_grain_percentiles
from reports.pptx_renderer import (
    render_pptx, PERCENTILE_TITLE, PERCENTILE_DEFINITION, MAX_DATA_ROWS,
)

HIER = [{"key": "sample", "label": "Part Number", "value": ""},
        {"key": "lot", "label": "Lot", "value": ""}]


def _img(iid, part, lot, diams_um, ppu=8.0, order=0):
    grains = [{"diameter_um": d, "area_um2": np.pi * (d / 2) ** 2,
               "diameter_px": d * ppu, "area_px": np.pi * (d * ppu / 2) ** 2} for d in diams_um]
    return ImageSummary(id=iid, image_path="", order=order, px_per_um=ppu, has_calibration=True,
                        grain_count=len(grains), grains=grains,
                        levels={"sample": part, "lot": lot})


def _model(images, units="um"):
    secs = [Section(id="cover", type="cover", title="Cover", order=0),
            Section(id="ov", type="overview_table", title="Overview", order=1)]
    return ReportModel(title="T", sections=secs, images=images, hierarchy=HIER, units=units)


def _text(slide):
    return "\n".join(r.text for sh in slide.shapes if sh.has_text_frame
                     for p in sh.text_frame.paragraphs for r in p.runs)


def _pct_slides(path):
    prs = Presentation(path)
    return prs, [s for s in prs.slides if _text(s).startswith(PERCENTILE_TITLE)]


def _table(slide):
    return next(sh for sh in slide.shapes if sh.has_table).table


def _rows(table):
    return [[table.cell(r, c).text for c in range(len(table.columns))]
            for r in range(1, len(table.rows))]


def _synthetic(n_parts=3, n_lots=2, seed=0):
    rng = np.random.default_rng(seed)
    imgs, k = [], 0
    for p in range(n_parts):
        for l in range(n_lots):
            for _ in range(2):
                imgs.append(_img(f"i{k}", f"P{p}", f"L{l}", rng.uniform(0.5, 5, 40), order=k))
                k += 1
    return imgs


def test_percentiles_pooled_match_numpy(tmp_path):
    a = _img("a", "P", "L", [1.0, 2.0, 3.0])
    b = _img("b", "P", "L", list(np.linspace(4, 40, 37)))
    out = str(tmp_path / "d.pptx")
    render_pptx(_model([a, b]), out)
    _, slides = _pct_slides(out)
    assert len(slides) == 1
    row = _rows(_table(slides[0]))[0]
    pool = np.array([g["diameter_um"] for g in a.grains + b.grains])
    p = np.percentile(pool, [10, 50, 90])
    assert row[2] == str(len(pool))
    assert [float(x) for x in row[3:6]] == [round(float(v), 2) for v in p]
    med_area = float(np.median([g["area_um2"] for g in a.grains + b.grains]))
    assert float(row[6]) == round(med_area, 2)
    # differs from an average of per-image percentiles
    avg = np.mean([np.percentile([g["diameter_um"] for g in i.grains], 50) for i in (a, b)])
    assert abs(avg - p[1]) > 0.01


def test_one_row_per_part_and_lot_and_definition(tmp_path):
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(_synthetic(3, 3)), out)
    prs, slides = _pct_slides(out)
    assert len(slides) == 1
    rows = _rows(_table(slides[0]))
    assert [(r[0], r[1]) for r in rows] == [(f"P{p}", f"L{l}") for p in range(3) for l in range(3)]
    assert PERCENTILE_DEFINITION in _text(slides[0])
    assert "D10: 10 % of grains are smaller than this size." in PERCENTILE_DEFINITION


def test_placed_after_summary_before_data_tables(tmp_path):
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(_synthetic(2, 2)), out)
    prs = Presentation(out)
    heads = [_text(s).split("\n")[0] for s in prs.slides]
    i = heads.index(PERCENTILE_TITLE)
    assert heads[i - 1].startswith("Grain Size Summary")
    assert heads[i + 1] == "P0"


def test_pagination_repeats_header_and_continued_title(tmp_path):
    imgs = _synthetic(n_parts=1, n_lots=MAX_DATA_ROWS + 1)
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(imgs), out)
    _, slides = _pct_slides(out)
    assert len(slides) == 2
    assert _text(slides[0]).startswith(f"{PERCENTILE_TITLE} (continued 1/2)")
    assert _text(slides[1]).startswith(f"{PERCENTILE_TITLE} (continued 2/2)")
    t0, t1 = _table(slides[0]), _table(slides[1])
    assert len(t0.rows) - 1 == MAX_DATA_ROWS and len(t1.rows) - 1 == 1
    assert [t0.cell(0, c).text for c in range(7)] == [t1.cell(0, c).text for c in range(7)]


def test_empty_lot_shows_dash(tmp_path):
    imgs = [_img("a", "P", "L1", [1, 2, 3]), _img("b", "P", "L2", [], order=1)]
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(imgs), out)
    _, slides = _pct_slides(out)
    row = _rows(_table(slides[0]))[1]
    assert row[2] == "0"
    assert row[3:] == ["–"] * 4
    assert "nan" not in _text(slides[0]).lower()


def test_unit_switch_only_rescales(tmp_path):
    imgs = _synthetic(1, 2)
    res = {}
    for u in ("um", "nm"):
        out = str(tmp_path / f"{u}.pptx")
        render_pptx(_model(imgs, units=u), out)
        _, slides = _pct_slides(out)
        tb = _table(slides[0])
        res[u] = (_rows(tb), [tb.cell(0, c).text for c in range(7)])
    assert "(µm)" in res["um"][1][3] and "(nm)" in res["nm"][1][3]
    assert "(µm²)" in res["um"][1][6] and "(nm²)" in res["nm"][1][6]
    for ru, rn in zip(res["um"][0], res["nm"][0]):
        assert ru[2] == rn[2]
        for c in (3, 4, 5):
            assert abs(float(rn[c]) - 1000 * float(ru[c])) <= 1000 * 0.005 + 1e-9
        assert abs(float(rn[6]) - 1e6 * float(ru[6])) <= 1e6 * 0.005 + 1e-6


def test_forbidden_word_absent(tmp_path):
    out = str(tmp_path / "d.pptx")
    render_pptx(_model(_synthetic(2, 2)), out)
    _, slides = _pct_slides(out)
    txt = _text(slides[0]).lower()
    assert "requestor" not in txt and "requester" not in txt


def test_helper_empty():
    r = pooled_grain_percentiles([])
    assert r["n"] == 0 and r["d50"] is None


def _lots_model(n_lots, part="P"):
    return _model([_img(f"i{k}", part, f"L{k}", [1.0 + k, 2.0 + k, 3.0 + k], order=k)
                   for k in range(n_lots)])


def test_pagination_boundaries(tmp_path):
    for n, expected in ((MAX_DATA_ROWS, [14]), (28, [14, 14]), (29, [14, 14, 1])):
        out = str(tmp_path / f"d{n}.pptx")
        render_pptx(_lots_model(n), out)
        _, slides = _pct_slides(out)
        assert [len(_table(s).rows) - 1 for s in slides] == expected
        if len(expected) > 1:
            assert _text(slides[-1]).startswith(f"{PERCENTILE_TITLE} (continued {len(expected)}/{len(expected)})")


def test_same_lot_name_under_two_parts_separate_pools(tmp_path):
    a = _img("a", "P1", "L1", [1.0, 2.0, 3.0])
    b = _img("b", "P2", "L1", [10.0, 20.0, 30.0], order=1)
    out = str(tmp_path / "d.pptx")
    render_pptx(_model([a, b]), out)
    _, slides = _pct_slides(out)
    rows = _rows(_table(slides[0]))
    assert [(r[0], r[1], r[2]) for r in rows] == [("P1", "L1", "3"), ("P2", "L1", "3")]
    assert float(rows[0][4]) == 2.0 and float(rows[1][4]) == 20.0


def test_excluded_images_not_pooled(tmp_path):
    a = _img("a", "P", "L", [1.0, 2.0, 3.0])
    b = _img("b", "P", "L", [100.0, 200.0], order=1)
    b.include = False
    out = str(tmp_path / "d.pptx")
    render_pptx(_model([a, b]), out)
    _, slides = _pct_slides(out)
    row = _rows(_table(slides[0]))[0]
    assert row[2] == "3" and float(row[4]) == 2.0


def test_nonfinite_grains_skipped():
    im = _img("a", "P", "L", [1.0, 2.0, 3.0])
    im.grains.append({"diameter_um": float("nan"), "area_um2": 1.0})
    im.grains.append({"diameter_um": 5.0, "area_um2": float("inf")})
    r = pooled_grain_percentiles([im])
    assert r["n"] == 3 and r["d50"] == 2.0


def test_uncalibrated_px_path(tmp_path):
    im = _img("a", "P", "L", [1.0, 2.0, 3.0], ppu=10.0)
    im.has_calibration = False
    im.px_per_um = 0.0
    out = str(tmp_path / "d.pptx")
    render_pptx(_model([im]), out)
    _, slides = _pct_slides(out)
    tb = _table(slides[0])
    assert "(px)" in tb.cell(0, 3).text and "(px²)" in tb.cell(0, 6).text
    assert float(_rows(tb)[0][4]) == 20.0  # 2.0 um * 10 px/um


def test_small_values_never_print_as_zero(tmp_path):
    from reports.pptx_renderer import _fmt_adaptive, _fmt_mean_sd
    assert _fmt_adaptive(0.004) == "0.00400"
    assert _fmt_adaptive(0.5) == "0.500" and _fmt_adaptive(12.345) == "12.35"
    assert _fmt_adaptive(0) == "0.00"
    assert _fmt_mean_sd(0.004, 0.0012, "µm²") == "0.00400 ± 0.00120 µm²"
    im = _img("a", "P", "L", [0.05, 0.06, 0.07])
    out = str(tmp_path / "d.pptx")
    render_pptx(_model([im]), out)
    _, slides = _pct_slides(out)
    row = _rows(_table(slides[0]))[0]
    assert all(float(x) > 0 for x in row[3:])


def test_layout_definition_wraps_and_table_clears_footer(tmp_path):
    from pptx.util import Emu
    out = str(tmp_path / "d.pptx")
    render_pptx(_lots_model(MAX_DATA_ROWS), out)
    prs, slides = _pct_slides(out)
    s = slides[0]
    box = next(sh for sh in s.shapes if sh.has_text_frame and PERCENTILE_DEFINITION in sh.text_frame.text)
    tbl = next(sh for sh in s.shapes if sh.has_table)
    assert box.text_frame.word_wrap
    assert Emu(box.top + box.height).inches <= Emu(tbl.top).inches
    assert Emu(tbl.top + tbl.height).inches <= prs.slide_height.inches - 0.32
