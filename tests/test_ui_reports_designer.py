"""
Reports page — the in-app report designer (REP-05/06/07 UI, UI-09).

Build a report from an analysed session, edit it (title, caption, disabled
section, image order, custom text, grain annotation), export XLSX + PPTX and
read the files back; reload report.json after recreating the window; the
"results changed" banner + Refresh numbers merge; excluding a grain from the
report goes through the app's manual exclusion so Review agrees.
"""
from pathlib import Path

import pytest

pytest.importorskip("pytestqt")

from data.models import AppSettings  # noqa: E402
from data.settings import save_settings  # noqa: E402
from tests.ui_shell_helpers import confirm_setup, make_session  # noqa: E402

TIMEOUT = 90000


@pytest.fixture
def env(tmp_path, monkeypatch, qapp):
    from ui.design.theme import apply_theme, set_reduced_motion
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path / "appdata"))
    root = tmp_path / "ws"
    save_settings(AppSettings(workspace_root=str(root), operator="Tester", theme="dark"))
    set_reduced_motion(True)
    apply_theme(qapp, "dark")
    yield root
    set_reduced_motion(False)


def _open_shell(qtbot, path: Path):
    from ui.app_shell import AppShell
    from ui.app_state import AppState
    shell = AppShell(AppState(), probe_device=False)
    qtbot.addWidget(shell)
    shell.resize(1500, 950)
    shell.show()
    shell.open_session(path)
    qtbot.waitUntil(lambda: shell.state.session is not None, timeout=15000)
    qtbot.waitUntil(lambda: not shell.reports.is_busy(), timeout=15000)
    return shell


def _analyse_all(shell, qtbot):
    st = shell.state
    shell.analyze.params.set_mode("threshold")
    confirm_setup(shell, qtbot)            # UX-02: scan area + scale confirmed first
    with qtbot.waitSignal(shell.analyze.queue.queue_finished, timeout=TIMEOUT):
        shell.analyze_all()
    qtbot.waitUntil(lambda: all(im.status == "done" and im.result is not None
                                for im in st.images()) and not st.is_filtering(),
                    timeout=TIMEOUT)


def _idle(shell, qtbot):
    st, rp = shell.state, shell.reports
    qtbot.waitUntil(lambda: not st.is_filtering() and not rp.is_busy(), timeout=TIMEOUT)
    qtbot.wait(450)   # stale-check debounce
    qtbot.waitUntil(lambda: not st.is_filtering() and not rp.is_busy(), timeout=TIMEOUT)


def _build(shell, qtbot):
    rp = shell.reports
    shell.go("reports")
    rp.build_from_session()
    qtbot.waitUntil(lambda: rp.model is not None and not rp.is_busy(), timeout=TIMEOUT)
    return rp


def _export(rp, qtbot, kinds):
    with qtbot.waitSignal(rp.exported, timeout=TIMEOUT) as blocker:
        rp.export(kinds)
    return blocker.args[0]


@pytest.fixture
def analysed(env, qtbot):
    path = make_session(env, 3, label="Run", black=True)
    shell = _open_shell(qtbot, path)
    _analyse_all(shell, qtbot)
    shell.state.flush()
    yield shell, path
    shell.close()


# ----------------------------------------------------------------------------

def test_reports_page_states(env, qtbot):
    path = make_session(env, 2, label="Fresh")
    shell = _open_shell(qtbot, path)
    rp = shell.reports
    shell.go("reports")
    assert rp.stack.currentWidget() is rp.empty_results     # nothing analysed yet
    _analyse_all(shell, qtbot)
    _idle(shell, qtbot)
    assert rp.stack.currentWidget() is rp.empty_build
    shell.state.close_session()
    assert rp.stack.currentWidget() is rp.empty_session
    shell.close()


def test_build_edit_export_roundtrip(analysed, qtbot, tmp_path):
    import openpyxl
    from pptx import Presentation
    shell, path = analysed
    rp = _build(shell, qtbot)
    m = rp.model
    assert len(m.images) == 3 and all(i.include for i in m.images)
    assert m.operator == "Tester"
    assert m.images[0].sample_id == "S-1" and m.images[0].lot_number == "L-1"
    assert (path / "report_assets").is_dir()
    assert rp.stack.currentWidget() is rp.designer

    # overview preview = one row per included image + combined row
    rp.select(("section", "overview_table"))
    assert rp.current_preview().rows() == 3

    # edit title (inspector), exclude an image, disable a section, reorder images
    rp.inspector.title.setText("Alloy 718 — Lot L-1")
    rp.inspector.title.textEdited.emit("Alloy 718 — Lot L-1")
    names = [Path(i.image_path).name for i in sorted(m.images, key=lambda i: i.order)]
    ids = [i.id for i in sorted(m.images, key=lambda i: i.order)]
    rp.set_image_included(ids[1], False)
    rp.set_section_enabled("combined_distribution", False)
    rp.move_image(ids[2], 0)                       # third image becomes first
    order = [i.id for i in sorted(rp.model.images, key=lambda i: i.order)]
    assert order == [ids[2], ids[0], ids[1]]
    assert rp.current_preview().rows() == 2        # overview follows the include flags

    # caption via the image preview's editor + a grain annotation
    rp.select(("image", ids[2]))
    prev = rp.current_preview()
    prev.caption.setText("Rim location, etched")
    gm = prev.gmodel
    row = next(r for r, x in enumerate(gm.rows) if x["kept"])
    gm.setData(gm.index(row, 7), "twin boundary")
    gid = gm.rows[row]["id"]
    # custom text section
    rp.select(("section", "overview_table"))
    rp.add_text_section()
    sec = next(s for s in rp.model.sections if s.type == "custom_text")
    rp.rename_section(sec.id, "Conclusions")
    sec.payload["body"] = "Meets ASTM E112 acceptance."
    rp.edited("text")

    xlsx, pptx = _export(rp, qtbot, ["xlsx", "pptx"])
    assert Path(xlsx).parent == path / "exports" and Path(pptx).parent == path / "exports"

    wb = openpyxl.load_workbook(xlsx)
    sheets = wb.sheetnames
    assert sheets[0] == "Overview" and "Summary Charts" not in sheets
    raw = [i for i, n in enumerate(sheets) if n.startswith("Raw - ")]
    assert raw and raw == list(range(len(sheets) - len(raw), len(sheets)))   # raw data last
    img_sheets = [n for n in sheets if n.startswith("Img ")]
    assert len(img_sheets) == 2
    assert Path(names[2]).stem in img_sheets[0] and Path(names[0]).stem in img_sheets[1]
    ov = wb["Overview"]
    assert ov["A1"].value == "Alloy 718 — Lot L-1"
    rows = list(ov.iter_rows(min_row=5, values_only=True))
    assert any(r[0] == "Lot" and r[1] == "Fields (n)" for r in rows)   # INN-27 lot block
    data_rows = [r for r in rows if isinstance(r[0], int) and r[1]
                 and r[1] != "Combined (all images)"]
    assert len(data_rows) == 2                         # overview rows = included images
    assert any(r[1] == "Combined (all images)" for r in ov.iter_rows(min_row=5, values_only=True))
    first_img = wb[img_sheets[0]]
    texts = " ".join(str(c.value) for row in first_img.iter_rows() for c in row if c.value)
    assert "Rim location, etched" in texts and f"#{gid}: twin boundary" in texts

    prs = Presentation(pptx)
    all_text = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                for s in prs.slides]
    assert "Alloy 718 — Lot L-1" in all_text[0]
    assert not any("Distribution" in t and "Combined" in t for t in all_text)
    img_slides = [i for i, t in enumerate(all_text) if "Image 1:" in t or "Image 2:" in t]
    assert len(img_slides) == 2 and "Rim location, etched" in all_text[img_slides[0]]
    concl = [i for i, t in enumerate(all_text) if "Conclusions" in t]
    assert concl and "Meets ASTM E112 acceptance." in all_text[concl[0]]
    assert concl[0] == 2          # right after the executive summary (charts disabled)
    assert "Appendix" in all_text[-1]
    pages = [t.split()[-1] for t in all_text]
    assert pages == [str(i) for i in range(1, len(prs.slides) + 1)]   # footers renumbered

    # export recorded in the report
    hist = rp.model.metadata["exports"]
    assert {h["file"] for h in hist} == {Path(xlsx).name, Path(pptx).name}
    assert all(h["operator"] == "Tester" for h in hist)


def test_report_persists_across_restart(analysed, qtbot):
    from data.session_io import load_session
    shell, path = analysed
    rp = _build(shell, qtbot)
    rp.inspector.org.setText("Metallography Lab")
    rp.inspector.org.textEdited.emit("Metallography Lab")
    first = sorted(rp.model.images, key=lambda i: i.order)
    rp.select(("image", first[0].id))
    rp.current_preview().caption.setText("Kept after restart")
    rp.set_section_enabled("parameters", False)
    rp.move_image(first[0].id, 2)
    shell.close()                                   # flush via closeEvent
    rep = load_session(path).report
    assert rep["organization"] == "Metallography Lab"

    shell2 = _open_shell(qtbot, path)
    rp2 = shell2.reports
    qtbot.waitUntil(lambda: rp2.model is not None, timeout=15000)
    m = rp2.model
    assert m.organization == "Metallography Lab"
    assert not m.get_section("parameters").enabled
    img = next(i for i in m.images if i.id == first[0].id)
    assert img.caption == "Kept after restart"
    assert img.order == 3
    shell2.go("reports")
    assert rp2.stack.currentWidget() is rp2.designer
    assert rp2.inspector.org.text() == "Metallography Lab"
    assert not rp2.is_stale()
    shell2.close()


def test_results_changed_banner_and_refresh_keeps_edits(analysed, qtbot):
    shell, path = analysed
    st = shell.state
    rp = _build(shell, qtbot)
    img0 = sorted(rp.model.images, key=lambda i: i.order)[0]
    rp.select(("image", img0.id))
    rp.current_preview().caption.setText("My caption")
    rp.inspector.bins_area.setValue(7)
    before = img0.grain_count
    # an edit made elsewhere (Review page) changes the numbers
    doc = rp.image_doc(img0)
    st.set_current_image(doc.uid)
    victims = [g.grain_id for g in doc.result.grains[:3]]
    assert st.delete_grains(doc.uid, victims)
    _idle(shell, qtbot)
    assert rp.is_stale() and rp.banner.isVisible()
    assert rp.model.images[0].grain_count == before        # not silently changed
    rp.refresh_numbers()
    qtbot.waitUntil(lambda: not rp.is_busy(), timeout=TIMEOUT)
    _idle(shell, qtbot)
    img = next(i for i in rp.model.images if i.id == img0.id)
    assert img.grain_count == before - 3
    assert img.caption == "My caption" and rp.model.bins["area"] == 7
    assert not rp.is_stale() and not rp.banner.isVisible()


def test_excluding_grain_from_report_matches_review(analysed, qtbot):
    shell, path = analysed
    st = shell.state
    rp = _build(shell, qtbot)
    img0 = sorted(rp.model.images, key=lambda i: i.order)[0]
    doc = rp.image_doc(img0)
    before = img0.grain_count
    rp.select(("image", img0.id))
    prev = rp.current_preview()
    gm = prev.gmodel
    row = next(r for r, x in enumerate(gm.rows) if x["kept"])
    gid = gm.rows[row]["id"]
    from PySide6.QtCore import Qt
    assert gm.setData(gm.index(row, 0), Qt.Unchecked, Qt.CheckStateRole)
    _idle(shell, qtbot)
    qtbot.waitUntil(lambda: next(i for i in rp.model.images if i.id == img0.id).grain_count
                    == before - 1, timeout=TIMEOUT)
    assert not rp.is_stale()                          # own edit: followed silently
    assert gid in doc.manual and doc.result.grain_count == before - 1
    # Review's comparison table shows the same count
    shell.go("review")
    shell.review._fill_comparison()
    r = [i for i in range(shell.review.cmp.rowCount())
         if shell.review.cmp.item(i, 0).text() == doc.display_name][0]
    assert shell.review.cmp.item(r, 2).text() == str(before - 1)
    # undo (Ctrl+Z) puts it back everywhere
    st.undo_stack.undo()
    _idle(shell, qtbot)
    qtbot.waitUntil(lambda: next(i for i in rp.model.images if i.id == img0.id).grain_count
                    == before, timeout=TIMEOUT)
    # re-tick from the report restores a hand-removed grain
    rp.set_grain_included(img0.id, gid, False)
    _idle(shell, qtbot)
    assert rp.set_grain_included(img0.id, gid, True)
    _idle(shell, qtbot)
    assert gid not in doc.manual


def test_menu_export_uses_report_pipeline(analysed, qtbot):
    import openpyxl
    shell, path = analysed
    rp = shell.reports
    assert rp.model is None
    with qtbot.waitSignal(rp.exported, timeout=TIMEOUT) as blocker:
        shell.export_all_excel()                      # Ctrl+E — builds the report on demand
    (xlsx,) = blocker.args[0]
    assert Path(xlsx).parent == path / "exports"
    assert openpyxl.load_workbook(xlsx).sheetnames[0] == "Overview"
    with qtbot.waitSignal(rp.exported, timeout=TIMEOUT) as blocker:
        shell.export_current_excel()                  # Ctrl+Shift+E — only the current image
    wb = openpyxl.load_workbook(blocker.args[0][0])
    assert len([n for n in wb.sheetnames if n.startswith("Img ")]) == 1


def test_validation_shows_missing_image_with_hint(analysed, qtbot):
    shell, path = analysed
    rp = _build(shell, qtbot)
    rp.model.images[0].image_path = str(path / "images" / "gone.png")
    probs = rp.validate()
    assert any(sev == "danger" and "missing" in text and "untick" in hint
               for sev, text, hint in probs)
    assert rp.issues_badge.isVisible()


def test_images_tab_stays_responsive_with_many_images(env, qtbot, monkeypatch):
    """UX-12: selecting the Images outline group / stepping through images
    must never block the GUI thread, even when every image decode is slow
    (large real-world SEM files) and there are many images in the report.

    Profiling this test is what found the actual freeze: ``QHeaderView.
    ResizeToContents`` on the per-grain table re-measures every cell in
    every column on every layout pass, so opening an image was
    O(grains x columns) — fine for the 3-image fixtures elsewhere, a real
    hang for a real SEM session. Fixed column widths made it O(1); the
    image-pixmap decode was also moved off the GUI thread for good measure
    (real SEM files are large)."""
    import time
    from ui import workers as ui_workers

    path = make_session(env, 50, label="Big")
    shell = _open_shell(qtbot, path)
    _analyse_all(shell, qtbot)
    rp = _build(shell, qtbot)
    assert len(rp.model.images) == 50

    # Simulate slow disk / large SEM files: every image decode sleeps.
    real_read = ui_workers.read_image

    def slow_read(p):
        time.sleep(0.05)
        return real_read(p)

    monkeypatch.setattr(ui_workers, "read_image", slow_read)
    rp._pix.clear()

    t0 = time.perf_counter()
    rp.select(("section", "images"))
    for img in rp.model.images[:20]:
        ta = time.perf_counter()
        rp.select(("image", img.id))
        assert time.perf_counter() - ta < 0.2, f"selecting image {img.id} blocked the GUI thread"
    elapsed = time.perf_counter() - t0
    assert elapsed < 2.0, f"opening the Images tab blocked the GUI thread for {elapsed:.2f}s"

    # The slow decodes still complete (off the GUI thread) and get cached.
    qtbot.waitUntil(lambda: rp.is_pixmaps_ready(rp.model.images[0]), timeout=10000)
    shell.close()


def test_custom_palette_saved_locally_and_listed_next_to_builtins(analysed, qtbot, monkeypatch):
    """UX-15: creating a custom palette persists it to AppSettings (local,
    survives restart), applies it to the current report, and lists it in
    the Palette combo next to the 4 built-ins on future reports."""
    shell, path = analysed
    rp = _build(shell, qtbot)
    insp = rp.inspector

    # Only the 4 built-ins + a separator + the trailing "New custom
    # palette..." entry so far (no custom palettes saved yet).
    from reports.charts import PALETTES
    assert insp.palette.count() == len(PALETTES) + 2

    class _FakeSignal:
        def __init__(self):
            self._cb = None

        def connect(self, cb):
            self._cb = cb

    class _FakeDialog:
        def __init__(self, *a, **kw):
            self.submitted = _FakeSignal()

        def exec(self):
            self._cb_result = self.submitted._cb("Lab blue",
                                                  ["#111111", "#222222", "#333333"])
            return 1

    monkeypatch.setattr("ui.dialogs.custom_palette_dialog.CustomPaletteDialog", _FakeDialog)
    insp._open_new_custom_palette_dialog()

    from reports.charts import derive_custom_palette
    expected = derive_custom_palette(["#111111", "#222222", "#333333"], "Lab blue")
    assert rp.model.theme == "custom:custom-1"
    assert rp.model.custom_palette == expected

    # Saved locally (AppSettings), not just on the model.
    assert shell.state.settings.custom_palettes == [
        {"id": "custom-1", "name": "Lab blue", "colors": ["#111111", "#222222", "#333333"]}]
    from data.settings import load_settings
    on_disk = load_settings(shell.state.settings_path)
    assert on_disk.custom_palettes == shell.state.settings.custom_palettes

    # Listed in the combo, selected, next to the 4 built-ins.
    assert insp.palette.count() == len(PALETTES) + 4   # 2 separators + 1 custom + "New..."
    assert insp.palette.currentData() == "custom:custom-1"
    assert insp.palette.currentText() == "Lab blue"

    # Selecting it again on a fresh inspector reload still applies it.
    rp.inspector.load()
    assert rp.inspector.palette.currentData() == "custom:custom-1"
    shell.close()


def test_legacy_ui_modules_are_gone():
    import importlib.util
    for mod in ("ui.main_window", "ui.settings_panel", "ui.results_panel",
                "ui.analysis_progress_dialog"):
        assert importlib.util.find_spec(mod) is None, mod
    root = Path(__file__).resolve().parents[1]
    for py in (root / "ui").rglob("*.py"):
        assert "excel_export" not in py.read_text(encoding="utf-8"), py
